"""Testes da feature de Linhagem do Lakehouse (integração Skill-LineageFabric):
parsing de notebooks, expansão transitiva, trilha SharePoint, persistência em
state.py e o grafo isolado (Mapa Isolado)."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import pytest

from dlctl.core import lineage_graph
from dlctl.core import state as state_db
from dlctl.core.global_scope import _extract_tables_from_sql, read_excel_pipeline_edges
from dlctl.core.dashboard_lineage import (
    _is_physical_table_name,
    _load_scan_dataflows,
    build_dashboard_lineage,
)
from dlctl.core.table_lineage_graph import build_mapped_table_dependency_graph
from dlctl.generators.lineage_generator import (
    build_sharepoint_dependency_trail,
    build_workspace_inventory,
    dependency_rows_from_dataframe,
    expand_transitive_lineage,
    export_lineage_excel,
    find_latest_lineage_excel,
    generate_lineage_artifacts,
    process_lakehouse_dev,
)
from dlctl.core.project_scan import _dashboard_table_names, _dashboards_prontos


def _write_silver_notebook(root: Path, domain: str) -> None:
    nb_dir = root / domain / "silver_notebooks" / "nb_silver"
    nb_dir.mkdir(parents=True, exist_ok=True)
    (nb_dir / "notebook-content.py").write_text(
        """
WORKSPACE_NAME = "LAKEHOUSE-DEV"
TARGET_LAKEHOUSE_LOGICAL = "LH_SUPRIMENTOS"
TARGET_SCHEMA = "silver"
SILVER_TABLE = "DW_RECEIPT_ORDER_CF"
WRITE_MODE = "overwrite"
BRONZE_SOURCES = [
    {"lakehouse": "LH_SUPRIMENTOS", "schema": "bronze", "table": "RCV_SHIPMENT_HEADERS"},
    {"lakehouse": "LH_SUPRIMENTOS", "schema": "bronze", "table": "RCV_SHIPMENT_LINES"},
]
""",
        encoding="utf-8",
    )


def _write_gold_notebook(root: Path, domain: str) -> None:
    nb_dir = root / domain / "gold_notebooks" / "nb_gold"
    nb_dir.mkdir(parents=True, exist_ok=True)
    (nb_dir / "notebook-content.py").write_text(
        """
WORKSPACE_NAME = "LAKEHOUSE-DEV"
# G1_HEADER_METADATA
GOLD_LAKEHOUSE_LOGICAL = "LH_SUPRIMENTOS"
GOLD_SCHEMA = "gold"
GOLD_TABLE = "PR_RECEIPT_ORDER"
DEPENDENCIES = [
    {"layer": "silver", "lakehouse": "LH_SUPRIMENTOS", "schema": "silver", "table": "DW_RECEIPT_ORDER_CF"},
]
""",
        encoding="utf-8",
    )


@pytest.fixture()
def lakehouse_dev_fixture(tmp_path):
    root = tmp_path / "lakehouse-dev"
    _write_silver_notebook(root, "SUPRIMENTOS")
    _write_gold_notebook(root, "SUPRIMENTOS")
    return root


def test_process_lakehouse_dev_parses_silver_and_gold(lakehouse_dev_fixture):
    result = process_lakehouse_dev(str(lakehouse_dev_fixture))
    linhagem = result["linhagem"]
    assert len(linhagem) == 3  # 2 bronze->silver + 1 silver->gold

    silver_rows = [r for r in linhagem if r["CAMADA_DESTINO"] == "silver"]
    assert {r["TABELA_ORIGEM"] for r in silver_rows} == {"RCV_SHIPMENT_HEADERS", "RCV_SHIPMENT_LINES"}
    assert all(r["TABELA_DESTINO"] == "DW_RECEIPT_ORDER_CF" for r in silver_rows)

    gold_rows = [r for r in linhagem if r["CAMADA_DESTINO"] == "gold"]
    assert len(gold_rows) == 1
    assert gold_rows[0]["TABELA_ORIGEM"] == "DW_RECEIPT_ORDER_CF"
    assert gold_rows[0]["TABELA_DESTINO"] == "PR_RECEIPT_ORDER"


def test_process_lakehouse_dev_empty_dir_returns_empty(tmp_path):
    result = process_lakehouse_dev(str(tmp_path / "does-not-exist"))
    assert result == {"tabelas": [], "linhagem": []}


def test_extract_tables_from_power_query_sql_decodes_structural_escapes():
    content = (
        "from fssupri.pr_purchase_order_val#(lf)) "
        "join fssupri.pr_requisition#(cr,lf) on 1 = 1"
    )

    assert _extract_tables_from_sql(content) == {
        "PR_PURCHASE_ORDER_VAL",
        "PR_REQUISITION",
    }


def test_load_scan_dataflows_prefers_decoded_mashup_document(tmp_path):
    scan_root = tmp_path / "scan"
    scan_root.mkdir()
    document = (
        'section Section1; shared Query = Oracle.Database("DB", '
        '[Query = "select * from fssupri.pr_purchase_order_val#(lf) '
        'join fssupri.pr_requisition on 1 = 1"]);'
    )
    payload = {
        "name": "DATAFLOW_SUPPLY_FUSION_02",
        "pbi:mashup": {"document": document},
    }
    (scan_root / "dataflow.json").write_text(json.dumps(payload), encoding="utf-8")

    loaded = _load_scan_dataflows(str(scan_root))

    assert loaded["DATAFLOW_SUPPLY_FUSION_02"] == document
    assert _extract_tables_from_sql(loaded["DATAFLOW_SUPPLY_FUSION_02"]) == {
        "PR_PURCHASE_ORDER_VAL",
        "PR_REQUISITION",
    }


def test_read_excel_pipeline_edges_preserves_links_when_layer_is_empty(tmp_path):
    import openpyxl

    sharedpoint_root = tmp_path / "sharedpoint"
    sharedpoint_root.mkdir()
    workbook_path = sharedpoint_root / "Projeto Lakehouse - Tabelas e Pipelines.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Tabelas"
    worksheet.cell(1, 7).value = "(0)\nSource"
    worksheet.cell(1, 10).value = "(2)\nCamada Bronze"
    worksheet.cell(1, 16).value = "(4)\nCamada Silver"
    worksheet.cell(1, 24).value = "(6)\nCamada Gold"
    worksheet.cell(2, 10).value = "BRONZE_TABLE"
    worksheet.cell(2, 24).value = "DM_GOLD_TABLE"
    worksheet.cell(3, 10).value = "BRONZE_TABLE"
    worksheet.cell(3, 16).value = "DW_SILVER_TABLE"
    worksheet.cell(3, 24).value = "DM_GOLD_TABLE"
    workbook.save(workbook_path)

    assert read_excel_pipeline_edges(str(sharedpoint_root)) == {
        ("BRONZE_TABLE", "DM_GOLD_TABLE"),
        ("BRONZE_TABLE", "DW_SILVER_TABLE"),
        ("DW_SILVER_TABLE", "DM_GOLD_TABLE"),
    }


def test_mapped_table_graph_follows_view_alias_and_deduplicates_sources(tmp_path):
    sharedpoint_root = tmp_path / "sharedpoint"
    stage = sharedpoint_root / "1 - Oracle ERP" / "6. Camada Gold"
    stage.mkdir(parents=True)
    other_root = tmp_path / "other_sources"
    other_root.mkdir()
    (stage / "DM_GOLD_TABLE.sql").write_text(
        "SELECT * FROM DW_SILVER_TABLE", encoding="utf-8"
    )
    (stage / "VW_DW_SILVER_TABLE.vw").write_text(
        "CREATE VIEW VW_DW_SILVER_TABLE AS SELECT * FROM BRONZE_TABLE",
        encoding="utf-8",
    )
    (other_root / "DW_SILVER_TABLE.vw").write_text(
        "CREATE VIEW DW_SILVER_TABLE AS SELECT * FROM BRONZE_TABLE",
        encoding="utf-8",
    )

    graph = build_mapped_table_dependency_graph(
        str(sharedpoint_root), {"DM_GOLD_TABLE"}
    )

    assert graph == {
        "DM_GOLD_TABLE": {"DW_SILVER_TABLE"},
        "DW_SILVER_TABLE": {"BRONZE_TABLE"},
    }


def test_dashboard_lineage_classifies_sharepoint_and_ignores_spaced_labels(tmp_path):
    workspaces_root = tmp_path / "Workspaces"
    workspaces_root.mkdir()
    payload = {
        "datasourceInstances": [{
            "datasourceType": "SharePointList",
            "connectionDetails": {"sharePointSiteUrl": "https://contoso.sharepoint.com/sites/data"},
            "datasourceId": "sp-1",
        }],
        "workspaces": [{
            "id": "ws-1", "name": "Workspace",
            "reports": [{"name": "Dashboard", "datasetId": "ds-1"}],
            "datasets": [{
                "id": "ds-1", "name": "Dataset",
                "tables": [{"name": "LINHAS DE RC"}, {"name": "DM_REAL_TABLE"}],
                "datasourceUsages": [{"datasourceInstanceId": "sp-1"}],
            }],
            "dataflows": [],
        }],
    }
    (workspaces_root / "workspace.json").write_text(json.dumps(payload), encoding="utf-8")

    result = build_dashboard_lineage(str(workspaces_root), str(tmp_path / "scan"), str(tmp_path / "sharedpoint"))

    assert _is_physical_table_name("DM_REAL_TABLE")
    assert not _is_physical_table_name("LINHAS DE RC")
    assert {row["tabela"] for row in result["dashboard_rows"] if row["tabela"]} == {"DM_REAL_TABLE"}
    model_row = next(row for row in result["dashboard_rows"] if row["tabela"] == "DM_REAL_TABLE")
    assert model_row["tipo_evidencia"] == "NOME_MODELO_CANDIDATO"
    assert model_row["confianca"] == "Baixa"
    assert model_row["tipo_objeto"] == "Entidade do modelo semântico"
    assert result["summary"]["total_referencias_alta_confianca"] == 0
    assert result["summary"]["total_candidatas_modelo_baixa_confianca"] == 1
    assert result["summary"]["qualidade_por_status"] == {"CANDIDATO_BAIXA_CONFIANCA": 1}
    assert result["dashboard_quality_rows"][0]["total_referencias"] == 1
    assert result["summary"]["total_fontes_sharepoint_bronze"] == 1
    assert result["sharepoint_rows"][0]["observacao"].startswith("Fonte SharePoint")


def test_dashboard_lineage_counts_same_name_reports_by_canonical_id(tmp_path):
    workspaces_root = tmp_path / "Workspaces"
    workspaces_root.mkdir()
    payload = {
        "datasourceInstances": [],
        "workspaces": [{
            "id": "ws-1", "name": "Workspace",
            "reports": [
                {"id": "rp-1", "name": "Mesmo nome", "datasetId": "ds-1"},
                {"id": "rp-2", "name": "Mesmo nome", "datasetId": "ds-2"},
            ],
            "datasets": [
                {"id": "ds-1", "name": "Dataset 1", "tables": [{"name": "DM_A"}]},
                {"id": "ds-2", "name": "Dataset 2", "tables": [{"name": "DM_B"}]},
            ],
            "dataflows": [],
        }],
    }
    (workspaces_root / "workspace.json").write_text(json.dumps(payload), encoding="utf-8")

    result = build_dashboard_lineage(
        str(workspaces_root), str(tmp_path / "scan"), str(tmp_path / "sharedpoint")
    )

    assert result["summary"]["total_dashboards"] == 2
    assert {row["report_id"] for row in result["dashboard_rows"]} == {"rp-1", "rp-2"}
    assert len(result["dashboard_quality_rows"]) == 2


def test_dashboard_lineage_counts_sharepoint_reports_by_canonical_id(tmp_path):
    workspaces_root = tmp_path / "Workspaces"
    workspaces_root.mkdir()
    payload = {
        "datasourceInstances": [{
            "datasourceType": "SharePointList",
            "connectionDetails": {"sharePointSiteUrl": "https://contoso.sharepoint.com/sites/data"},
            "datasourceId": "sp-1",
        }],
        "workspaces": [{
            "id": "ws-1", "name": "Workspace",
            "reports": [
                {"id": "rp-1", "name": "Mesmo nome", "datasetId": "ds-1"},
                {"id": "rp-2", "name": "Mesmo nome", "datasetId": "ds-2"},
            ],
            "datasets": [
                {"id": "ds-1", "name": "Dataset 1", "tables": [],
                 "datasourceUsages": [{"datasourceInstanceId": "sp-1"}]},
                {"id": "ds-2", "name": "Dataset 2", "tables": [],
                 "datasourceUsages": [{"datasourceInstanceId": "sp-1"}]},
            ],
            "dataflows": [],
        }],
    }
    (workspaces_root / "workspace.json").write_text(json.dumps(payload), encoding="utf-8")

    result = build_dashboard_lineage(
        str(workspaces_root), str(tmp_path / "scan"), str(tmp_path / "sharedpoint")
    )

    assert result["summary"]["total_dashboards_com_sharepoint"] == 2


def test_dashboard_lineage_exposes_client_layer_conflicts(tmp_path):
    import openpyxl

    workspaces_root = tmp_path / "Workspaces"
    workspaces_root.mkdir()
    (workspaces_root / "workspace.json").write_text(
        json.dumps({"datasourceInstances": [], "workspaces": []}), encoding="utf-8"
    )
    sharedpoint_root = tmp_path / "sharedpoint"
    sharedpoint_root.mkdir()
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Tabelas"
    worksheet.cell(1, 16).value = "(4)\nCamada Silver"
    worksheet.cell(1, 24).value = "(6)\nCamada Gold"
    worksheet.cell(2, 16).value = "PR_ORDEM_SERVICO"
    worksheet.cell(3, 24).value = "PR_ORDEM_SERVICO"
    workbook.save(sharedpoint_root / "Projeto Lakehouse - Tabelas e Pipelines.xlsx")

    result = build_dashboard_lineage(
        str(workspaces_root), str(tmp_path / "scan"), str(sharedpoint_root)
    )

    assert result["summary"]["total_conflitos_camada_excel_cliente"] == 1
    assert result["client_layer_conflicts"] == [{
        "tabela": "PR_ORDEM_SERVICO", "camadas": "Gold, Silver"
    }]


def test_expand_transitive_lineage_bridges_bronze_to_gold(lakehouse_dev_fixture):
    linhagem = process_lakehouse_dev(str(lakehouse_dev_fixture))["linhagem"]
    expanded = expand_transitive_lineage(linhagem)

    transitive_rows = [r for r in expanded if r.get("_is_transitive")]
    assert len(transitive_rows) == 2  # os 2 bronzes agora também apontam direto para gold
    for row in transitive_rows:
        assert row["CAMADA_ORIGEM"] == "bronze"
        assert row["CAMADA_DESTINO"] == "gold"
        assert row["TABELA_DESTINO"] == "PR_RECEIPT_ORDER"
        assert "[TRANSITIVA]" in row["OBSERVACAO"]

    # relações diretas continuam intactas
    direct_rows = [r for r in expanded if not r.get("_is_transitive")]
    assert len(direct_rows) == 3


def test_expand_transitive_lineage_empty_input_is_empty():
    assert expand_transitive_lineage([]) == []


def _write_sharepoint_workspace_json(root: Path) -> None:
    root.mkdir(parents=True, exist_ok=True)
    payload = {
        "datasourceInstances": [],
        "workspaces": [
            {
                "id": "ws-1",
                "name": "SUPRIMENTOS WS",
                "datasets": [
                    {
                        "id": "ds-1",
                        "name": "Dataset Recebimento",
                        "tables": [
                            {
                                "name": "TabelaRecebimento",
                                "source": [{
                                    "expression": (
                                        'let\n Source = SharePoint.Files('
                                        '"https://contoso.sharepoint.com/sites/Suprimentos")\nin Source'
                                    )
                                }],
                            }
                        ],
                    }
                ],
                "reports": [
                    {"id": "rpt-1", "name": "Dashboard Recebimento", "datasetId": "ds-1"},
                    {"id": "rpt-2", "name": "Dashboard Orfao", "datasetId": "ds-missing"},
                ],
            }
        ],
    }
    (root / "workspace1.json").write_text(json.dumps(payload), encoding="utf-8")


def test_build_sharepoint_dependency_trail_detects_direct_and_report_chain(tmp_path):
    workspaces_root = tmp_path / "Workspaces"
    _write_sharepoint_workspace_json(workspaces_root)

    rows = build_sharepoint_dependency_trail(str(workspaces_root))

    direct = [r for r in rows if r["chain_depth"] == 0 and r["report_name"] == ""]
    assert len(direct) == 1
    assert direct[0]["exists_check"] == "existe"
    assert direct[0]["sharepoint_reference"] == "https://contoso.sharepoint.com/sites/Suprimentos"

    via_report = [r for r in rows if r["report_name"] == "Dashboard Recebimento"]
    assert len(via_report) == 1
    assert via_report[0]["exists_check"] == "existe"
    assert via_report[0]["chain_depth"] == 1

    orphan = [r for r in rows if r["report_name"] == "Dashboard Orfao"]
    assert len(orphan) == 1
    assert orphan[0]["exists_check"] == "nao_encontrado"


def test_build_sharepoint_dependency_trail_missing_root_returns_empty(tmp_path):
    assert build_sharepoint_dependency_trail(str(tmp_path / "missing")) == []


def test_build_workspace_inventory_flattens_all_item_types(tmp_path):
    workspaces_root = tmp_path / "Workspaces"
    _write_sharepoint_workspace_json(workspaces_root)

    rows = build_workspace_inventory(str(workspaces_root))
    by_type = {}
    for row in rows:
        by_type.setdefault(row["item_type"], []).append(row)

    assert len(by_type["Workspace"]) == 1
    assert by_type["Workspace"][0]["item_name"] == "SUPRIMENTOS WS"
    assert by_type["Workspace"][0]["workspace_id"] == "ws-1"

    assert len(by_type["Dataset"]) == 1
    assert by_type["Dataset"][0]["item_name"] == "Dataset Recebimento"
    assert "table_count=1" in by_type["Dataset"][0]["detail"]

    assert len(by_type["Dataset Table"]) == 1
    assert by_type["Dataset Table"][0]["item_name"] == "TabelaRecebimento"
    assert by_type["Dataset Table"][0]["parent_name"] == "Dataset Recebimento"
    assert "dataset_id=ds-1" in by_type["Dataset Table"][0]["detail"]

    assert len(by_type["Report"]) == 2
    report_names = {r["item_name"] for r in by_type["Report"]}
    assert report_names == {"Dashboard Recebimento", "Dashboard Orfao"}
    assert all(r["workspace"] == "SUPRIMENTOS WS" for r in rows)


def test_build_workspace_inventory_missing_root_returns_empty(tmp_path):
    assert build_workspace_inventory(str(tmp_path / "missing")) == []


def test_dashboard_table_mapping_is_scoped_by_workspace_and_dataset_id():
    rows = [
        {"workspace": "WS A", "workspace_id": "ws-a", "item_type": "Dataset",
         "item_name": "Modelo Compartilhado", "item_id": "ds-a", "parent_name": "", "detail": ""},
        {"workspace": "WS A", "workspace_id": "ws-a", "item_type": "Dataset Table",
         "item_name": "DM_A", "item_id": "", "parent_name": "Modelo Compartilhado",
         "detail": "dataset_id=ds-a, storage_mode=Import"},
        {"workspace": "WS A", "workspace_id": "ws-a", "item_type": "Report",
         "item_name": "Relatorio A", "item_id": "report-a", "parent_name": "ds-a", "detail": ""},
        {"workspace": "WS B", "workspace_id": "ws-b", "item_type": "Dataset",
         "item_name": "Modelo Compartilhado", "item_id": "ds-b", "parent_name": "", "detail": ""},
        {"workspace": "WS B", "workspace_id": "ws-b", "item_type": "Dataset Table",
         "item_name": "DM_B", "item_id": "", "parent_name": "Modelo Compartilhado",
         "detail": "dataset_id=ds-b, storage_mode=DirectLake"},
        {"workspace": "WS B", "workspace_id": "ws-b", "item_type": "Report",
         "item_name": "Relatorio B", "item_id": "report-b", "parent_name": "ds-b", "detail": ""},
    ]

    mapping = _dashboard_table_names(rows)

    assert mapping == {"report-a": {"DM_A"}, "report-b": {"DM_B"}}


def test_dashboard_ready_requires_all_known_gold_dependencies():
    dashboard_tables = {
        "complete": {"DM_A", "DM_B"},
        "partial": {"DM_A", "DM_MISSING"},
        "unknown": set(),
    }

    assert _dashboards_prontos(dashboard_tables, {"DM_A", "DM_B"}) == (3, 1)


def test_generate_lineage_artifacts_persists_workspace_inventory(tmp_profile, lakehouse_dev_fixture, tmp_path):
    workspaces_root = tmp_path / "Workspaces"
    _write_sharepoint_workspace_json(workspaces_root)

    result = generate_lineage_artifacts(
        tmp_profile, lakehouse_dev_input=str(lakehouse_dev_fixture), workspaces_input=str(workspaces_root),
    )

    assert result["workspace_item_rows"] > 0
    items = state_db.get_lineage_workspace_items(tmp_profile, result["batch_id"])
    assert len(items) == result["workspace_item_rows"]
    assert any(i.item_type == "Report" and i.item_name == "Dashboard Recebimento" for i in items)

    batch = state_db.latest_lineage_batch(tmp_profile)
    assert batch.workspace_item_rows == result["workspace_item_rows"]


def test_generate_lineage_artifacts_persists_and_exports(tmp_profile, lakehouse_dev_fixture, tmp_path):
    workspaces_root = tmp_path / "Workspaces"
    _write_sharepoint_workspace_json(workspaces_root)

    result = generate_lineage_artifacts(
        tmp_profile,
        lakehouse_dev_input=str(lakehouse_dev_fixture),
        workspaces_input=str(workspaces_root),
    )

    assert result["dependency_rows"] == 5  # 3 diretas + 2 transitivas
    assert result["sharepoint_rows"] == 3
    assert result["excel_path"] and Path(result["excel_path"]).exists()

    batch = state_db.latest_lineage_batch(tmp_profile)
    assert batch is not None
    assert batch.batch_id == result["batch_id"]
    assert batch.status == "success"

    dependencies = state_db.get_lineage_dependencies(tmp_profile, result["batch_id"])
    assert len(dependencies) == 5
    assert any(dep.is_transitive for dep in dependencies)

    sharepoint_rows = state_db.get_lineage_sharepoint(tmp_profile, result["batch_id"])
    assert {row.exists_check for row in sharepoint_rows} == {"existe", "nao_encontrado"}


def test_generate_lineage_artifacts_without_workspaces_input_skips_sharepoint(tmp_profile, lakehouse_dev_fixture):
    result = generate_lineage_artifacts(
        tmp_profile, lakehouse_dev_input=str(lakehouse_dev_fixture), workspaces_input=None,
    )
    assert result["sharepoint_rows"] == 0


def test_lineage_graph_isolate_upstream_and_downstream(lakehouse_dev_fixture):
    linhagem = process_lakehouse_dev(str(lakehouse_dev_fixture))["linhagem"]
    expanded = expand_transitive_lineage(linhagem)

    dependency_dicts = [{
        "source_layer": r["CAMADA_ORIGEM"], "source_lakehouse": r["LAKEHOUSE_ORIGEM"],
        "source_schema": r["SCHEMA_ORIGEM"], "source_table": r["TABELA_ORIGEM"],
        "target_layer": r["CAMADA_DESTINO"], "target_lakehouse": r["LAKEHOUSE_DESTINO"],
        "target_schema": r["SCHEMA_DESTINO"], "target_table": r["TABELA_DESTINO"],
        "dependency_type": r["TIPO_DEPENDENCIA"], "domain": r.get("Domínio", ""), "note": r.get("OBSERVACAO", ""),
    } for r in expanded]

    graph = lineage_graph.build_dependency_graph(dependency_dicts)

    gold_node = lineage_graph.layer_node_id("gold", "LH_SUPRIMENTOS", "gold", "PR_RECEIPT_ORDER")
    upstream = lineage_graph.isolate_nodes(graph, [gold_node], direction="Upstream")
    summary = lineage_graph.summarize_by_layer(upstream)
    assert "RCV_SHIPMENT_HEADERS" in summary["bronze"]
    assert "RCV_SHIPMENT_LINES" in summary["bronze"]
    assert "DW_RECEIPT_ORDER_CF" in summary["silver"]
    assert "PR_RECEIPT_ORDER" in summary["gold"]

    bronze_node = lineage_graph.layer_node_id("bronze", "LH_SUPRIMENTOS", "bronze", "RCV_SHIPMENT_HEADERS")
    downstream = lineage_graph.isolate_nodes(graph, [bronze_node], direction="Downstream")
    downstream_summary = lineage_graph.summarize_by_layer(downstream)
    assert "PR_RECEIPT_ORDER" in downstream_summary["gold"]


def test_lineage_graph_find_nodes_by_term(lakehouse_dev_fixture):
    linhagem = process_lakehouse_dev(str(lakehouse_dev_fixture))["linhagem"]
    dependency_dicts = [{
        "source_layer": r["CAMADA_ORIGEM"], "source_lakehouse": r["LAKEHOUSE_ORIGEM"],
        "source_schema": r["SCHEMA_ORIGEM"], "source_table": r["TABELA_ORIGEM"],
        "target_layer": r["CAMADA_DESTINO"], "target_lakehouse": r["LAKEHOUSE_DESTINO"],
        "target_schema": r["SCHEMA_DESTINO"], "target_table": r["TABELA_DESTINO"],
        "dependency_type": r["TIPO_DEPENDENCIA"],
    } for r in linhagem]
    graph = lineage_graph.build_dependency_graph(dependency_dicts)

    matches = lineage_graph.find_nodes_by_term(graph, "dw_receipt_order_cf")
    assert len(matches) == 1
    assert graph.nodes[matches[0]]["table"] == "DW_RECEIPT_ORDER_CF"


def test_dependency_rows_from_dataframe_builds_full_graph(lakehouse_dev_fixture):
    """A pagina Grafo Isolado le o Excel exportado; garante que o DataFrame
    reconstruido a partir dele produz o mesmo grafo que o batch original,
    incluindo deteccao de relacoes transitivas pela tag em OBSERVACAO."""
    linhagem = process_lakehouse_dev(str(lakehouse_dev_fixture))["linhagem"]
    expanded = expand_transitive_lineage(linhagem)

    from dlctl.generators.lineage_generator import COLS_LINEAGE_TABLES

    df = pd.DataFrame(expanded, columns=COLS_LINEAGE_TABLES)
    rows = dependency_rows_from_dataframe(df)

    assert len(rows) == len(expanded)
    transitive_count = sum(1 for r in rows if r["is_transitive"])
    assert transitive_count == 2

    graph = lineage_graph.build_dependency_graph(rows)
    gold_node = lineage_graph.layer_node_id("gold", "LH_SUPRIMENTOS", "gold", "PR_RECEIPT_ORDER")
    upstream = lineage_graph.isolate_nodes(graph, [gold_node], direction="Upstream")
    summary = lineage_graph.summarize_by_layer(upstream)
    assert "RCV_SHIPMENT_HEADERS" in summary["bronze"]
    assert "DW_RECEIPT_ORDER_CF" in summary["silver"]


def test_dependency_rows_from_dataframe_empty_returns_empty():
    from dlctl.generators.lineage_generator import COLS_LINEAGE_TABLES

    empty_df = pd.DataFrame(columns=COLS_LINEAGE_TABLES)
    assert dependency_rows_from_dataframe(empty_df) == []


def test_find_latest_lineage_excel_picks_most_recent(tmp_profile):
    lineage_dir = tmp_profile.paths.manifests_root / "lineage"
    lineage_dir.mkdir(parents=True, exist_ok=True)

    older = export_lineage_excel([], [], lineage_dir / "lineage_tables_batch_old.xlsx")
    import time
    time.sleep(0.05)
    newer = export_lineage_excel([], [], lineage_dir / "lineage_tables_batch_new.xlsx")

    found = find_latest_lineage_excel(tmp_profile)
    assert found == newer
    assert found != older


def test_find_latest_lineage_excel_missing_dir_returns_none(tmp_profile):
    assert find_latest_lineage_excel(tmp_profile) is None
