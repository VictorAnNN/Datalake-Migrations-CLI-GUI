"""dlctl.core.global_scope

Escopo TOTAL do projeto (quantas tabelas precisam ser migradas ao todo),
independente do que já foi feito em `input/lakehouse-dev`/`input/lakehouse-hml`
(isso é "existente"; este módulo calcula o "esperado" real do projeto).

Fonte da verdade: `input/sharedpoint/Projeto Lakehouse - Tabelas e
Pipelines.xlsx`, aba "Tabelas" — planilha mantida pelo cliente com uma linha
por pipeline (Source -> Bronze -> Silver -> Gold), com uma coluna "Camada X"
por camada (`(2)\\nCamada Bronze`, `(4)\\nCamada Silver`, `(6)\\nCamada
Gold`, e qualquer `Camada <Y>` adicional que o cliente venha a criar — as
colunas de camada são descobertas dinamicamente pelo cabeçalho, não fixas
por índice).

Problema conhecido: nem toda tabela que precisa ser migrada está mapeada
nessa planilha. Para cobrir isso, também vasculhamos os scripts `.sql`/`.prc`
dentro de cada pasta de sistema fonte em `input/sharedpoint/` (ex.:
`1 - Oracle ERP/`, `2 - Maximo/`, `3 - RM/`, ... cada uma com suas próprias
subpastas de estágio "1. Script Source x Bronze", "3. Script Bronze x
Silver", "4. Camada Silver", "5. Script Silver x Gold"), procurando tabelas
referenciadas via `FROM`/`JOIN` que não apareçam em nenhuma coluna de camada
do Excel. Essas tabelas "descobertas" são incluídas no total e classificadas
pela convenção de prefixo do cliente:
- `DW_`  -> Silver (Prata)
- `DM_` ou `PR_` -> Gold
- fora da regra -> Bronze
Se a tabela vier como view (`VW_`/`V_` na frente), o prefixo de view é
ignorado e a classificação segue o prefixo mapeado logo depois
(ex.: `VW_DW_PEDIDO` -> Silver, `V_PR_PEDIDO` -> Gold).
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Mesma paleta usada pelos relatórios do Supervisor (duplicada aqui, em vez de
# importada de dlctl.core.project_scan, para evitar import circular).
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
EVEN_FILL = PatternFill("solid", fgColor="DCE6F1")
ODD_FILL = PatternFill("solid", fgColor="FFFFFF")

EXCEL_RELATIVE_PATH = "Projeto Lakehouse - Tabelas e Pipelines.xlsx"
SHEET_NAME = "Tabelas"
DOMAIN_COLUMN_HEADER = "Domínio"

# Aba com o plano de entregáveis do projeto — onde o total de dashboards a
# migrar está declarado em texto livre (ex.: "Reapontamento dos 428
# dashboards" / "428 dashboards conectados à nova origem").
DELIVERABLES_SHEET_NAME = "Entregáveis LKH"
_DASHBOARD_COUNT_PATTERN = re.compile(r"([\d][\d.,]*)\s*dashboards?\b", re.IGNORECASE)

# Nomes de subpasta de estágio (dentro de CADA pasta de sistema fonte em
# input/sharedpoint/, ex.: "1 - Oracle ERP/1. Script Source x Bronze/")
# vasculhadas em busca de tabelas referenciadas via FROM/JOIN que não
# estejam em nenhuma coluna de camada do Excel.
SQL_SCAN_FOLDERS = [
    "1. Script Source x Bronze",
    "3. Script Bronze x Silver",
    "4. Camada Silver",
    "5. Script Silver x Gold",
]

# Modo "estendido" (segundo artefato do Supervisor): além das pastas/arquivos
# acima, também vasculha "2. Camada Bronze" e "6. Camada Gold" (onde vivem os
# .tab de criação de tabela) e os arquivos .prc/.tab/.vw/.dsx (procedures,
# DDLs de tabela/view Oracle e exports de job DataStage), que o modo padrão
# ignora. É uma superfície de busca maior — não substitui o artefato
# original, é comparado a ele na tela do Supervisor.
EXTENDED_SCAN_FOLDERS = SQL_SCAN_FOLDERS + ["2. Camada Bronze", "6. Camada Gold"]
DEFAULT_SCAN_EXTENSIONS = (".sql", ".prc")
EXTENDED_SCAN_EXTENSIONS = (".sql", ".prc", ".tab", ".vw", ".dsx")

# Fonte extra (modo estendido): pasta solta `input/other_sources/` (irmã de
# `input/sharedpoint/`), sem a estrutura "sistema fonte / pasta de estágio"
# — vasculhada por inteiro (rglob), com `.txt` a mais na lista de extensões
# porque alguns scripts lá foram salvos com nome tipo "X.prc.txt". Pode ter
# conteúdo duplicado do que já existe em input/sharedpoint; a deduplicação
# por nome de tabela (`discovered`/`known_tables`) já cobre isso.
OTHER_SOURCES_DIR_NAME = "other_sources"
OTHER_SOURCES_SCAN_EXTENSIONS = (".sql", ".prc", ".tab", ".vw", ".dsx", ".txt")

_IGNORED_TABLE_VALUES = {"", "NA", "N/A", "TBD", "-"}
_SQL_TABLE_STOPLIST = {"DUAL"}
_FROM_JOIN_PATTERN = re.compile(
    r"\b(?:FROM|JOIN)\s+([A-Za-z_][A-Za-z0-9_$#]*(?:\.[A-Za-z_][A-Za-z0-9_$#]*)?)",
    re.IGNORECASE,
)

# Nome do próprio objeto criado num .tab/.vw (ex.: "create table FSSUPRI.PR_X"
# ou "CREATE OR REPLACE FORCE VIEW FUSION.VW_X AS") — só usado no modo
# estendido, para capturar a tabela/view definida no arquivo, não apenas o
# que ela referencia via FROM/JOIN.
_CREATE_OBJECT_PATTERN = re.compile(
    r"\bCREATE\s+(?:OR\s+REPLACE\s+)?(?:FORCE\s+)?(?:TABLE|VIEW)\s+"
    r"[A-Za-z_][A-Za-z0-9_$#]*\.([A-Za-z_][A-Za-z0-9_$#]*)",
    re.IGNORECASE,
)

# Tabelas referenciadas em jobs DataStage exportados (.dsx), ex.:
# TableDef "Database\og05des\FUSION.VW_DW_PURCHASE_ORDER_APPROVAL_HIST" —
# só usado no modo estendido.
_DSX_TABLEDEF_PATTERN = re.compile(
    r'TableDef\s+"[^"\\]*\\[A-Za-z_][A-Za-z0-9_$#]*\.([A-Za-z_][A-Za-z0-9_$#]*)"',
    re.IGNORECASE,
)

# Nome de uma CTE ("WITH nome AS (...)", ou o próximo bloco encadeado por
# vírgula "..., outro_nome AS (...)") — é só um apelido de subquery em
# memória, não uma tabela física, então nunca deve ser tratado como tabela
# descoberta mesmo quando reaparece num FROM/JOIN mais abaixo no mesmo script.
_CTE_ALIAS_PATTERN = re.compile(r"\b([A-Za-z_][A-Za-z0-9_$#]*)\s+AS\s*\(", re.IGNORECASE)


def _norm(name: str) -> str:
    return str(name or "").strip().upper()


def _classify_by_prefix(table_name: str) -> str:
    """Regra de prefixo do cliente, usada só para tabelas descobertas via SQL
    (que não têm uma coluna de camada própria no Excel): DW_ -> Silver, DM_/
    PR_ -> Gold, qualquer outra coisa -> Bronze. O prefixo TEM que vir
    seguido de "_" (ex.: "DW_X" conta, "DWQTE_X" não é DW_ e cai em Bronze).

    Se o nome vier prefixado com `VW_`/`V_` (view sobre uma tabela DW/DM/PR),
    o prefixo de view é descartado antes de checar a convenção, já que a
    classificação real segue o prefixo mapeado que vem depois."""
    name = _norm(table_name)
    if name.startswith("VW_"):
        name = name[len("VW_"):]
    elif name.startswith("V_"):
        name = name[len("V_"):]
    if name.startswith("DW_"):
        return "Silver"
    if name.startswith("DM_") or name.startswith("PR_"):
        return "Gold"
    return "Bronze"


def _find_camada_columns(header_row: tuple) -> dict[str, int]:
    """Descobre dinamicamente as colunas 'Camada X' pelo cabeçalho (ex.:
    '(2)\\nCamada Bronze' -> {'Bronze': 9}) — não depende de índice fixo,
    então cobre automaticamente qualquer camada nova que o cliente crie."""
    camada_cols: dict[str, int] = {}
    for i, h in enumerate(header_row):
        if not h:
            continue
        text = str(h).replace("\n", " ").strip()
        lower = text.lower()
        if "camada" in lower:
            idx = lower.index("camada")
            nome = text[idx + len("camada"):].strip()
            if nome:
                camada_cols[nome] = i
    return camada_cols


def _find_domain_column(header_row: tuple) -> Optional[int]:
    for i, h in enumerate(header_row):
        if h and str(h).strip() == DOMAIN_COLUMN_HEADER:
            return i
    return None


def read_dashboard_target(sharedpoint_input: str, sheet_name: str = DELIVERABLES_SHEET_NAME) -> Optional[dict]:
    """Lê o total de dashboards a migrar declarado em texto livre na aba do
    plano de entreg\u00e1veis (ex.: \"428 dashboards conectados \u00e0 nova origem\") \u2014
    n\u00e3o \u00e9 contado, \u00e9 o n\u00famero que o pr\u00f3prio cliente definiu como escopo."""
    excel_path = Path(sharedpoint_input) / EXCEL_RELATIVE_PATH
    if not excel_path.exists():
        return None
    wb = openpyxl.load_workbook(str(excel_path), read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        return None
    ws = wb[sheet_name]
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if not cell or not isinstance(cell, str):
                continue
            match = _DASHBOARD_COUNT_PATTERN.search(cell)
            if not match:
                continue
            raw = match.group(1).replace(".", "").replace(",", "")
            try:
                total = int(raw)
            except ValueError:
                continue
            return {"total": total, "fonte_sheet": sheet_name, "fonte_texto": cell.strip()}
    return None


def read_excel_tables(sharedpoint_input: str) -> dict:
    """Lê a aba 'Tabelas' do Excel e retorna, por camada (Bronze/Silver/Gold/
    quaisquer outras), o conjunto de nomes de tabela únicos + o detalhe por
    linha (tabela, camada, domínio)."""
    excel_path = Path(sharedpoint_input) / EXCEL_RELATIVE_PATH
    result = {
        "excel_found": False, "excel_path": str(excel_path),
        "camada_names": [], "tables_by_camada": {}, "rows": [],
    }
    if not excel_path.exists():
        return result

    wb = openpyxl.load_workbook(str(excel_path), read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        return result
    ws = wb[SHEET_NAME]

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return result

    camada_cols = _find_camada_columns(header_row)
    domain_col = _find_domain_column(header_row)
    if not camada_cols:
        return result

    tables_by_camada: dict[str, set[str]] = {nome: set() for nome in camada_cols}
    rows: list[dict] = []
    for row in rows_iter:
        domain = str(row[domain_col]).strip() if domain_col is not None and row[domain_col] else ""
        for nome, col_idx in camada_cols.items():
            value = row[col_idx] if col_idx < len(row) else None
            if not value or not isinstance(value, str):
                continue
            table_name = value.strip()
            if _norm(table_name) in _IGNORED_TABLE_VALUES:
                continue
            norm_name = _norm(table_name)
            if norm_name in tables_by_camada[nome]:
                continue
            tables_by_camada[nome].add(norm_name)
            rows.append({
                "tabela": norm_name, "camada": nome, "dominio": domain,
                "origem": "Excel", "detalhe": "",
            })

    result.update({
        "excel_found": True,
        "camada_names": list(camada_cols.keys()),
        "tables_by_camada": tables_by_camada,
        "rows": rows,
    })
    return result


_IDENTIFIER_LIKE_PATTERN = re.compile(r"^[A-Z0-9_\-\.]+$")


def read_excel_pipeline_edges(sharedpoint_input: str) -> set[tuple[str, str]]:
    """Lê a aba 'Tabelas' de novo, mas capturando a LIGAÇÃO entre as colunas
    de camada de cada linha (mesma linha = mesmo pipeline do cliente): se
    duas colunas de camada consecutivas (na ordem em que aparecem no
    cabeçalho, ex. Bronze -> Silver -> Gold) tiverem valor preenchido na
    mesma linha, vira uma aresta tabela_origem -> tabela_destino. É a
    linhagem Bronze/Silver/Gold já mapeada manualmente pelo cliente na
    planilha — mais confiável que qualquer heurística sobre os scripts
    (usada como fonte extra de linhagem em `dlctl.core.table_lineage_graph`).

    Valores que não parecem nome de tabela/objeto (com espaço, frase livre
    tipo "Criar Tabela SA") são descartados — só entram identificadores
    (letras/números/`_`/`-`/`.`)."""
    excel_path = Path(sharedpoint_input) / EXCEL_RELATIVE_PATH
    edges: set[tuple[str, str]] = set()
    if not excel_path.exists():
        return edges

    wb = openpyxl.load_workbook(str(excel_path), read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        return edges
    ws = wb[SHEET_NAME]

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return edges

    camada_cols = _find_camada_columns(header_row)
    if len(camada_cols) < 2:
        return edges
    ordered_col_indexes = [col_idx for _, col_idx in sorted(camada_cols.items(), key=lambda kv: kv[1])]

    def _clean(value) -> str:
        if not value or not isinstance(value, str):
            return ""
        text = _norm(value)
        if text in _IGNORED_TABLE_VALUES or not _IDENTIFIER_LIKE_PATTERN.match(text):
            return ""
        return text

    for row in rows_iter:
        values = [
            _clean(row[col_idx] if col_idx < len(row) else None)
            for col_idx in ordered_col_indexes
        ]
        values = [value for value in values if value]
        for source, target in zip(values, values[1:]):
            if source and target and source != target:
                edges.add((source, target))
    return edges


def _extract_tables_from_sql(content: str, suffix: str = "") -> set[str]:
    """Extrai nomes de tabela de um arquivo de script. No modo padrão, só
    olha para FROM/JOIN (SQL/PL-SQL). No modo estendido (`suffix` de um
    arquivo .tab/.vw/.dsx), também captura o nome do próprio objeto criado
    (.tab/.vw) e tabelas referenciadas em jobs DataStage (.dsx).

    Nomes de CTE (bloco `WITH nome AS (...)`) nunca entram como tabela
    descoberta — são só apelidos de subquery em memória, mesmo quando
    reaparecem num FROM/JOIN mais abaixo no script (ex.: `LEFT JOIN
    DMSUBINV`, onde DMSUBINV é a CTE, não uma tabela física)."""
    # Exports do Power Query preservam escapes M dentro da SQL nativa
    # (por exemplo ``FROM schema.table#(lf)``). Se o texto bruto for
    # submetido ao regex, o ``#`` — válido em identificadores Oracle — é
    # anexado ao nome da tabela e gera falsos objetos como ``TABLE#``.
    # Decodificamos somente os escapes estruturais antes da análise; nomes
    # Oracle que realmente contenham ``#`` continuam aceitos normalmente.
    content = (
        content.replace("#(cr,lf)", "\n")
        .replace("#(cr)", "\n")
        .replace("#(lf)", "\n")
        .replace("#(tab)", "\t")
    )

    cte_aliases = {_norm(m.group(1)) for m in _CTE_ALIAS_PATTERN.finditer(content)}

    tables: set[str] = set()
    for match in _FROM_JOIN_PATTERN.finditer(content):
        raw = match.group(1)
        name = raw.split(".")[-1]
        name = _norm(name)
        if name and name not in _SQL_TABLE_STOPLIST and name not in cte_aliases:
            tables.add(name)

    suffix = suffix.lower()
    if suffix in (".tab", ".vw"):
        for match in _CREATE_OBJECT_PATTERN.finditer(content):
            name = _norm(match.group(1))
            if name and name not in _SQL_TABLE_STOPLIST:
                tables.add(name)
    elif suffix == ".dsx":
        for match in _DSX_TABLEDEF_PATTERN.finditer(content):
            name = _norm(match.group(1))
            if name and name not in _SQL_TABLE_STOPLIST:
                tables.add(name)
    return tables


def discover_tables_from_sql(sharedpoint_input: str, known_tables: set[str], extended: bool = False) -> list[dict]:
    """Vasculha, dentro de CADA pasta de sistema fonte em `input/sharedpoint/`
    (ex.: `1 - Oracle ERP/`, `2 - Maximo/`, ...), as subpastas de estágio em
    `SQL_SCAN_FOLDERS` (ex.: `1. Script Source x Bronze/**/*.sql`/`.prc`) por
    tabelas citadas via FROM/JOIN que ainda não estão em `known_tables`
    (união de tudo já encontrado no Excel) — essas são tabelas que precisam
    ser migradas mas não estão mapeadas na planilha.

    Com `extended=True` (segundo artefato do Supervisor), também vasculha
    `EXTENDED_SCAN_FOLDERS` e os arquivos `.tab`/`.vw`/`.dsx` além de
    `.sql`/`.prc` — cobre DDLs de tabela/view Oracle e jobs DataStage
    exportados que o modo padrão não olha. Além disso, também vasculha por
    inteiro a pasta solta `input/other_sources/` (irmã de `sharedpoint_input`,
    sem a estrutura de sistema fonte/est\u00e1gio), incluindo `.txt` — mesmo que
    tenha conteúdo duplicado do que já existe em sharedpoint, cada tabela só
    entra uma vez (dedup por nome, via `known_tables`/`discovered`)."""
    root = Path(sharedpoint_input)
    discovered: dict[str, dict] = {}
    if not root.exists():
        return []

    scan_folders = EXTENDED_SCAN_FOLDERS if extended else SQL_SCAN_FOLDERS
    scan_extensions = EXTENDED_SCAN_EXTENSIONS if extended else DEFAULT_SCAN_EXTENSIONS

    for domain_dir in sorted(p for p in root.iterdir() if p.is_dir()):
        for stage_folder_name in scan_folders:
            stage_root = domain_dir / stage_folder_name
            if not stage_root.exists():
                continue
            scanned_paths = []
            for ext in scan_extensions:
                scanned_paths.extend(stage_root.rglob(f"*{ext}"))
            for sql_path in scanned_paths:
                try:
                    content = sql_path.read_text(encoding="utf-8", errors="replace")
                except Exception:
                    continue

                for table_name in _extract_tables_from_sql(content, sql_path.suffix):
                    if table_name in known_tables or table_name in discovered:
                        continue
                    origem = "SQL (não estava no Excel)"
                    if sql_path.suffix.lower() in (".tab", ".vw", ".dsx"):
                        origem = f"{sql_path.suffix.lstrip('.').upper()} (não estava no Excel)"
                    discovered[table_name] = {
                        "tabela": table_name, "camada": _classify_by_prefix(table_name),
                        "dominio": domain_dir.name, "origem": origem,
                        "detalhe": f"{stage_folder_name} / {sql_path.name}",
                    }

    if extended:
        other_root = root.parent / OTHER_SOURCES_DIR_NAME
        if other_root.exists():
            scanned_paths = []
            for ext in OTHER_SOURCES_SCAN_EXTENSIONS:
                scanned_paths.extend(other_root.rglob(f"*{ext}"))
            for sql_path in scanned_paths:
                try:
                    content = sql_path.read_text(encoding="utf-8", errors="replace")
                except Exception:
                    continue

                for table_name in _extract_tables_from_sql(content, sql_path.suffix):
                    if table_name in known_tables or table_name in discovered:
                        continue
                    ext_label = sql_path.suffix.lstrip(".").upper() or "SQL"
                    discovered[table_name] = {
                        "tabela": table_name, "camada": _classify_by_prefix(table_name),
                        "dominio": "", "origem": f"{ext_label} (input/other_sources)",
                        "detalhe": f"other_sources / {sql_path.relative_to(other_root)}",
                    }
    return list(discovered.values())



def compute_global_scope(sharedpoint_input: str = "input/sharedpoint", extended: bool = False) -> dict:
    """Calcula o escopo TOTAL do projeto (quantas tabelas precisam ser
    migradas ao todo, por camada), a partir do Excel do cliente + descoberta
    via SQL de tabelas referenciadas que não estão mapeadas na planilha.

    Com `extended=True`, a descoberta também vasculha `.tab`/`.vw`/`.dsx`
    (além de `.sql`/`.prc`) e as pastas `2. Camada Bronze`/`6. Camada Gold` —
    ver `discover_tables_from_sql`."""
    excel_data = read_excel_tables(sharedpoint_input)
    if not excel_data["excel_found"]:
        return {
            "excel_found": False, "excel_path": excel_data["excel_path"],
            "camadas": [], "total_tabelas": 0, "rows": [], "descobertas_sql": [],
            "dashboard_target": None, "extended": extended,
        }

    known_tables: set[str] = set()
    for tables in excel_data["tables_by_camada"].values():
        known_tables |= tables

    descobertas_sql = discover_tables_from_sql(sharedpoint_input, known_tables, extended=extended)

    camada_counts: dict[str, int] = {
        nome: len(tables) for nome, tables in excel_data["tables_by_camada"].items()
    }
    for d in descobertas_sql:
        camada_counts[d["camada"]] = camada_counts.get(d["camada"], 0) + 1

    camadas = [
        {
            "camada": nome,
            "existente_excel": len(excel_data["tables_by_camada"].get(nome, set())),
            "descobertas_sql": sum(1 for d in descobertas_sql if d["camada"] == nome),
            "total": camada_counts.get(nome, 0),
        }
        for nome in sorted(camada_counts, key=lambda n: -camada_counts[n])
    ]

    return {
        "excel_found": True,
        "excel_path": excel_data["excel_path"],
        "camadas": camadas,
        "total_tabelas": sum(camada_counts.values()),
        "rows": excel_data["rows"] + descobertas_sql,
        "descobertas_sql": descobertas_sql,
        "dashboard_target": read_dashboard_target(sharedpoint_input),
        "extended": extended,
    }


def _autofit(ws_sheet) -> None:
    for col_idx, col in enumerate(ws_sheet.iter_cols(), start=1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for i, cell in enumerate(col):
            if i == 0:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
            else:
                cell.fill = EVEN_FILL if i % 2 == 0 else ODD_FILL
            try:
                max_len = max(max_len, len(str(cell.value)) if cell.value else 0)
            except Exception:
                pass
        ws_sheet.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 80)
    ws_sheet.freeze_panes = "A2"


def export_global_scope_report(scope_result: dict, output_dir: Path, batch_id: str) -> dict[str, str]:
    """Exporta o escopo total do projeto para `manifests/dashboard/`: um
    Excel com abas Resumo/Tabelas/Descobertas via SQL (esta última só para
    conferência discreta de gaps do Excel oficial).

    Se `scope_result["extended"]` for True (artefato do modo estendido —
    também vasculha .tab/.vw/.dsx), o artefato é salvo com o prefixo
    `escopo_estendido_` em vez de `escopo_total_`, para não se misturar com
    o histórico do artefato original."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    prefix = "escopo_estendido" if scope_result.get("extended") else "escopo_total"
    excel_path = output_dir / f"{prefix}_{batch_id}.xlsx"

    wb = openpyxl.Workbook()
    ws_resumo = wb.active
    ws_resumo.title = "Resumo"
    ws_resumo.append(["Camada", "Existente no Excel", "Descobertas via SQL", "Total"])
    for c in scope_result["camadas"]:
        ws_resumo.append([c["camada"], c["existente_excel"], c["descobertas_sql"], c["total"]])
    ws_resumo.append(["TOTAL GERAL", "", "", scope_result["total_tabelas"]])
    if scope_result.get("dashboard_target"):
        dt = scope_result["dashboard_target"]
        ws_resumo.append(["Dashboards/BI (meta declarada)", "", "", dt["total"]])
        ws_resumo.append([f"  fonte: aba '{dt['fonte_sheet']}' — \"{dt['fonte_texto']}\"", "", "", ""])
    _autofit(ws_resumo)

    ws_tabelas = wb.create_sheet("Tabelas")
    ws_tabelas.append(["Tabela", "Camada", "Domínio", "Origem", "Detalhe"])
    for r in scope_result["rows"]:
        ws_tabelas.append([r["tabela"], r["camada"], r.get("dominio", ""), r["origem"], r.get("detalhe", "")])
    _autofit(ws_tabelas)

    ws_gaps = wb.create_sheet("Descobertas via SQL")
    ws_gaps.append(["Tabela", "Camada (por prefixo)", "Domínio", "Encontrada em"])
    for d in scope_result["descobertas_sql"]:
        ws_gaps.append([d["tabela"], d["camada"], d.get("dominio", ""), d.get("detalhe", "")])
    _autofit(ws_gaps)

    wb.save(str(excel_path))

    # Snapshot em JSON (mesmo batch_id) para o dashboard recarregar o último
    # artefato gerado sem precisar reprocessar Excel/SQL a cada acesso.
    snapshot = dict(scope_result)
    snapshot["batch_id"] = batch_id
    snapshot["generated_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    snapshot["excel_path"] = str(excel_path)
    json_path = output_dir / f"{prefix}_{batch_id}.json"
    json_path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")

    return {"excel_path": str(excel_path), "json_path": str(json_path)}


def load_latest_scope_snapshot(output_dir: Path, extended: bool = False) -> Optional[dict]:
    """Carrega o snapshot (`escopo_total_*.json`, ou `escopo_estendido_*.json`
    se `extended=True`) mais recente já exportado em `output_dir`, para o
    dashboard sempre abrir com o último artefato gerado sem precisar
    reprocessar Excel/SQL — só recalcula quando o usuário pedir
    explicitamente (botão de gerar novo artefato)."""
    output_dir = Path(output_dir)
    if not output_dir.exists():
        return None
    prefix = "escopo_estendido" if extended else "escopo_total"
    candidates = sorted(output_dir.glob(f"{prefix}_*.json"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not candidates:
        return None
    try:
        return json.loads(candidates[0].read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


def list_scope_snapshots(output_dir: Path, extended: bool = False) -> list[dict]:
    """Lista todos os snapshots já gerados (mais recente primeiro) — usado
    pelo histórico de gerações do dashboard. Com `extended=True`, lista os
    snapshots do artefato estendido (`escopo_estendido_*.json`) em vez do
    original (`escopo_total_*.json`)."""
    output_dir = Path(output_dir)
    if not output_dir.exists():
        return []
    prefix = "escopo_estendido" if extended else "escopo_total"
    snapshots = []
    for p in sorted(output_dir.glob(f"{prefix}_*.json"), key=lambda p: p.stat().st_mtime, reverse=True):
        try:
            snapshots.append(json.loads(p.read_text(encoding="utf-8")))
        except (OSError, json.JSONDecodeError):
            continue
    return snapshots
