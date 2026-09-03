"""dlctl.core.dashboard_lineage

Descobre a linhagem de TODOS os dashboards/relatórios do tenant, cruzando:
- `input/Workspaces` (Fabric Scanner API): Report -> Dataset -> tabelas do
  modelo, e/ou -> Dataflow (via `upstreamDataflows`/`targetDataflowId`);
- `input/scan` (exports brutos de Dataflow, com o código M/Power Query
  embutido): tabelas físicas realmente referenciadas via FROM/JOIN dentro de
  cada query do Dataflow — muito mais confiável que o nome do modelo do
  Dataset, que costuma ser um nome amigável dado pelo desenvolvedor;
- `input/sharedpoint` (Excel "Projeto Lakehouse..." + scripts .sql/.prc/
  .tab/.vw/.dsx, via `compute_global_scope(extended=True)`): usado só como
  apoio para saber a camada/domínio de cada tabela física já conhecida
  (tabelas que não aparecem lá caem no fallback de classificação por
  prefixo — DW_ -> Silver, DM_/PR_ -> Gold, resto -> Bronze).

Além disso, a partir de cada tabela resolvida (geralmente Gold), rastreia
PRA TRÁS via `dlctl.core.table_lineage_graph` (dependência tabela -> tabela
extraída dos mesmos scripts .sql/.prc/.tab/.vw de `input/sharedpoint`) até
achar as tabelas Silver/Bronze que a alimentam — por isso o total de
tabelas/camadas cresce em relação ao que vem só de Workspaces/Scan.

As referências dos Dataflows, das M-queries internas e dos nomes das entidades
do modelo são combinadas. Nomes de entidade continuam best-effort, enquanto
`FROM`/`JOIN` extraídos das M-queries recebem evidência estática forte."""
from __future__ import annotations

import json
import re
from collections import Counter, deque
from pathlib import Path

import openpyxl

from dlctl.core.global_scope import (
    _autofit,
    _classify_by_prefix,
    _extract_tables_from_sql,
    _is_physical_table_name as _is_valid_physical_table_name,
    _norm,
    compute_global_scope,
    read_excel_tables,
)
from dlctl.core.table_lineage_graph import (
    build_mapped_table_dependency_graph,
    build_table_source_system_hints,
    build_table_dependency_graph,
    trace_upstream,
)
from dlctl.core.dashboard_lineage_html import export_dashboard_lineage_html
from dlctl.generators.lineage_generator import _load_json_files, _read_json

_TRAILING_NUMBER_SUFFIX = re.compile(r"\s+\d+$")
_GENERIC_MODEL_TABLE_NAMES = {
    "TASKS", "CALENDAR", "FEATURES", "MEASUREMENTS",
}


def _normalize_dataflow_key(name: str) -> str:
    """Normaliza o nome de um Dataflow/arquivo de scan para comparação (ex.:
    o Dataflow real se chama 'CONTABIL', mas o arquivo exportado em
    input/scan foi salvo como 'CONTABIL 1.json' — o ' 1' é descartado)."""
    return _TRAILING_NUMBER_SUFFIX.sub("", _norm(name))


def _is_physical_table_name(name: str) -> bool:
    """Aceita somente identificadores de objeto físico, não rótulos do modelo.

    Nomes como ``LINHAS DE RC`` e ``LEAD TIME`` são nomes amigáveis de
    tabelas do modelo Power BI, não nomes de tabela física usados no
    mapeamento. Espaços, portanto, desqualificam o valor nesta etapa.
    """
    normalized = _norm(name)
    return normalized not in _GENERIC_MODEL_TABLE_NAMES and _is_valid_physical_table_name(normalized)


def _reference_evidence(
    table_name: str,
    direct_origins: dict[str, str],
    documented_tables: set[str],
) -> tuple[str, str, str]:
    """Qualifica a evidência sem confundir nome do modelo com objeto físico.

    O Scanner chama de ``tables`` tanto entidades do modelo semântico quanto
    tabelas físicas. Uma referência extraída de M/SQL ou alcançada pelo
    mapping explícito é evidência estática forte de dependência. Um nome do
    modelo que também aparece no inventário é corroborado, mas nenhum desses
    sinais prova existência no motor. Os demais nomes permanecem candidatos.
    """
    table = _norm(table_name)
    direct_origin = direct_origins.get(table)
    if not direct_origin:
        return "REFERENCIA_UPSTREAM_ESTATICA", "Alta", "Tabela ou view"
    if direct_origin == "via Dataflow (M-query)":
        return "REFERENCIA_M_QUERY", "Alta", "Tabela ou view"
    if direct_origin == "via Dataset M-query":
        return "REFERENCIA_M_QUERY_MODELO", "Alta", "Tabela ou view"
    if table in documented_tables:
        return "NOME_MODELO_DOCUMENTADO", "Média", "Objeto documentado pelo cliente"
    return "NOME_MODELO_CANDIDATO", "Baixa", "Entidade do modelo semântico"


def _load_scan_dataflows(scan_input: str) -> dict[str, str]:
    """Carrega os JSONs de `input/scan` (exports de Dataflow com código M
    embutido) e indexa pelo nome normalizado -> documento M decodificado.

    O JSON bruto contém escapes do próprio JSON e escapes estruturais do
    Power Query (por exemplo, ``#(lf)``). Analisar o arquivo serializado pode
    concatenar o fim de um identificador com o token seguinte e também omitir
    referências reais. O conteúdo bruto fica apenas como fallback para scans
    legados que não tenham ``pbi:mashup.document``.
    """
    root = Path(scan_input)
    scans: dict[str, str] = {}
    if not root.exists():
        return scans
    for jf in sorted(root.glob("*.json")):
        try:
            content = jf.read_text(encoding="utf-8", errors="replace")
        except OSError:
            continue
        try:
            payload = json.loads(content)
            name = payload.get("name") or jf.stem
            mashup = payload.get("pbi:mashup") or {}
            document = mashup.get("document", "") if isinstance(mashup, dict) else ""
        except (json.JSONDecodeError, AttributeError):
            name = jf.stem
            document = ""
        scans[_normalize_dataflow_key(name)] = document or content
    return scans


def _build_end_to_end_mapping_rows(
    client_tables: set[str],
    client_layers_by_table: dict[str, set[str]],
    client_domains_by_table: dict[str, set[str]],
    dependencies_by_target: dict[str, set[str]],
    dashboard_endpoints: dict[tuple[str, str, str, str, str, str], dict[str, dict]],
    classify,
    client_sources_by_table: dict[str, set[str]] | None = None,
    script_sources_by_table: dict[str, set[str]] | None = None,
    rejected_rows: list[dict] | None = None,
) -> list[dict]:
    """Diagnostica cada tabela do mapping até um endpoint de dashboard.

    O grafo recebido usa ``destino -> origens``. Aqui ele é invertido e
    percorrido no sentido Bronze/Silver/Gold -> consumidor. O menor caminho
    é evidência estática; não comprova execução ou materialização no Fabric.
    """
    downstream: dict[str, set[str]] = {}
    for target, sources in dependencies_by_target.items():
        for source in sources:
            if source and target and source != target:
                downstream.setdefault(source, set()).add(target)

    endpoint_consumers: dict[str, list[tuple[tuple[str, str, str, str, str, str], dict]]] = {}
    for key, endpoints in dashboard_endpoints.items():
        for table, evidence in endpoints.items():
            endpoint_consumers.setdefault(table, []).append((key, evidence))

    confidence_rank = {"Alta": 0, "Média": 1, "Baixa": 2, "": 3}
    layer_rank = {"Bronze": 0, "Silver": 1, "Gold": 2}
    client_sources_by_table = client_sources_by_table or {}
    script_sources_by_table = script_sources_by_table or {}
    rows: list[dict] = []

    source_labels = {
        "ORACLE ERP": "Oracle ERP (Fusion)",
        "MAXIMO": "Máximo",
    }

    def _display_sources(values: set[str]) -> str:
        labels = {
            source_labels.get(value.strip().upper(), value.strip())
            for value in values if value and value.strip()
        }
        return ", ".join(sorted(labels, key=str.casefold))

    def _resolve_source_system(table: str) -> tuple[str, str]:
        direct = set(client_sources_by_table.get(table, set()))
        if direct:
            return _display_sources(direct), "Mapping do cliente (coluna Sistema)"
        scripted = set(script_sources_by_table.get(table, set()))
        if scripted:
            return _display_sources(scripted), "Sistema explícito no caminho/comentário do script"

        # Não há inferência por prefixo. Para linhas Gold/Silver sem Sistema,
        # herda-se somente a origem da dependência upstream explícita mais
        # próxima, mantendo a decisão auditável.
        visited = {table}
        queue = deque([(table, 0)])
        nearest_distance: int | None = None
        inherited: set[str] = set()
        while queue:
            current, distance = queue.popleft()
            if nearest_distance is not None and distance >= nearest_distance:
                continue
            for source in sorted(dependencies_by_target.get(current, set())):
                if source in visited:
                    continue
                visited.add(source)
                source_values = (
                    set(client_sources_by_table.get(source, set()))
                    or set(script_sources_by_table.get(source, set()))
                )
                if source_values:
                    nearest_distance = distance + 1
                    inherited.update(source_values)
                else:
                    queue.append((source, distance + 1))
        if inherited:
            return _display_sources(inherited), "Herdado da dependência upstream explícita mais próxima"
        return "Não identificado", "Sem Sistema no mapping ou dependência upstream qualificada"

    for mapped_table in sorted(client_tables):
        parent: dict[str, str | None] = {mapped_table: None}
        distance = {mapped_table: 0}
        queue = deque([mapped_table])
        while queue:
            current = queue.popleft()
            for target in sorted(downstream.get(current, set())):
                if target in parent:
                    continue
                parent[target] = current
                distance[target] = distance[current] + 1
                queue.append(target)

        reachable_endpoints = [table for table in parent if table in endpoint_consumers]
        consumers: list[tuple[tuple[str, str, str, str, str, str], str, dict]] = []
        for endpoint in reachable_endpoints:
            consumers.extend((key, endpoint, evidence) for key, evidence in endpoint_consumers[endpoint])
        consumers.sort(key=lambda item: (
            confidence_rank.get(item[2].get("confianca", ""), 3),
            distance[item[1]], item[0][1], item[0][3], item[1],
        ))

        declared_layers = sorted(client_layers_by_table.get(mapped_table, set()))
        domains = sorted(value for value in client_domains_by_table.get(mapped_table, set()) if value)
        source_system, source_evidence = _resolve_source_system(mapped_table)
        alerts = []
        if len(declared_layers) > 1:
            alerts.append(f"CONFLITO_CAMADA: {', '.join(declared_layers)}")
        base = {
            "tabela_mapeada": mapped_table,
            "camada_mapeada": ", ".join(declared_layers) or classify(mapped_table)[0],
            "dominio": ", ".join(domains),
            "sistema_origem": source_system,
            "evidencia_sistema_origem": source_evidence,
            "alertas": "; ".join(alerts),
        }

        if consumers:
            best_key, endpoint, evidence = consumers[0]
            path = []
            cursor: str | None = endpoint
            while cursor is not None:
                path.append(cursor)
                cursor = parent[cursor]
            path.reverse()
            canonical_consumers = {
                ((key[2],) if key[2] else (key[0], key[3], key[4]))
                for key, _, _ in consumers
            }
            dashboard_names = sorted({f"{key[1]} / {key[3]}" for key, _, _ in consumers})
            sample = "; ".join(dashboard_names[:5])
            if len(dashboard_names) > 5:
                sample += f"; +{len(dashboard_names) - 5}"

            confidence = evidence.get("confianca", "")
            if confidence == "Baixa":
                status, reaches = "CANDIDATO_BAIXA_CONFIANCA", "Candidato"
                observation = (
                    "O caminho chega somente a um nome de entidade do modelo semântico; "
                    "o nome físico ainda não foi corroborado por M-query, SQL ou mapping."
                )
                action = "Confirmar o nome físico no modelo/Dataflow antes de contabilizar como cobertura."
            elif len(path) == 1:
                status, reaches = "COMPLETO_DIRETO", "Sim"
                observation = "A própria tabela mapeada é referenciada diretamente pelo dataset/Dataflow do dashboard."
                action = "Validar existência, materialização e consulta no Fabric."
            else:
                status, reaches = "COMPLETO", "Sim"
                observation = "Cadeia estática confirmada até um endpoint consumido por dashboard, por nomes completos."
                action = "Validar execução/materialização no Fabric e o teste funcional do dashboard."
            if alerts:
                observation += " Há conflito de camada declarado no mapping do cliente."
            rows.append({
                **base,
                "status_fim_a_fim": status, "chega_dashboard": reaches,
                "etapa_alcancada": "Dashboard", "qtd_dashboards": len(canonical_consumers),
                "dashboards_amostra": sample, "tabela_endpoint": endpoint,
                "caminho_exemplo": " -> ".join(path),
                "workspace_id": best_key[0], "workspace": best_key[1],
                "report_id": best_key[2], "dashboard": best_key[3],
                "dataset_id": best_key[4], "dataset": best_key[5],
                "tipo_evidencia_endpoint": evidence.get("tipo_evidencia", ""),
                "confianca": confidence, "codigo_motivo": status,
                "observacao": observation, "acao_recomendada": action,
            })
            continue

        reached_tables = set(parent)
        reached_layer = max(
            (classify(table)[0] for table in reached_tables),
            key=lambda layer: layer_rank.get(layer, -1),
            default=classify(mapped_table)[0],
        )
        if mapped_table not in downstream:
            status = "ORFAO_SEM_ARESTA"
            observation = (
                "A tabela consta no mapping do cliente, mas nenhum script, view ou ligação do Excel "
                "fornecido indica um consumidor downstream."
            )
            action = "Localizar o script/pipeline consumidor ou confirmar descontinuação/fora de escopo."
            path_text = mapped_table
        else:
            status = "PARCIAL_SEM_DASHBOARD"
            furthest = max(
                reached_tables,
                key=lambda table: (layer_rank.get(classify(table)[0], -1), distance[table], table),
            )
            path = []
            cursor = furthest
            while cursor is not None:
                path.append(cursor)
                cursor = parent[cursor]
            path_text = " -> ".join(reversed(path))
            observation = (
                f"Há cadeia estática até a camada {reached_layer}, mas nenhum endpoint coincide com "
                "a linhagem de dashboard disponível."
            )
            if reached_layer == "Gold":
                action = "Verificar dataset/report consumidor, nome físico versus modelo e cobertura do Scanner API."
            elif reached_layer == "Silver":
                action = "Localizar a transformação Silver→Gold ou confirmar consumo direto pela Gold/dashboard."
            else:
                action = "Localizar a transformação Bronze→Silver/Gold."
        if alerts:
            observation += " Há conflito de camada declarado no mapping do cliente."
        rows.append({
            **base,
            "status_fim_a_fim": status, "chega_dashboard": "Não",
            "etapa_alcancada": reached_layer, "qtd_dashboards": 0,
            "dashboards_amostra": "", "tabela_endpoint": "",
            "caminho_exemplo": path_text,
            "workspace_id": "", "workspace": "", "report_id": "", "dashboard": "",
            "dataset_id": "", "dataset": "", "tipo_evidencia_endpoint": "",
            "confianca": "", "codigo_motivo": status,
            "observacao": observation, "acao_recomendada": action,
        })
    for rejected in sorted(rejected_rows or [], key=lambda row: (row["valor"], row["camada"])):
        value = rejected["valor"]
        if re.search(r"\.(?:XLSX?|CSV|PARQUET)$", value, re.IGNORECASE):
            status = "FONTE_ARQUIVO_NAO_TABELA"
            observation = "A entrada é um arquivo de origem, não um nome físico de tabela."
            action = "Registrar a ingestão Bronze e o nome da tabela física materializada a partir do arquivo."
        elif re.fullmatch(r"[A-Z_][A-Z0-9_$#-]*\.[A-Z_][A-Z0-9_$#-]*", value):
            status = "NOME_QUALIFICADO_NAO_NORMALIZADO"
            observation = "A entrada usa nome qualificado (schema.objeto), enquanto o grafo compara nomes físicos canônicos."
            action = "Separar schema e objeto; confirmar qual nome deve ser comparado no Lakehouse."
        elif re.fullmatch(r"[A-Z0-9_$#-]*_[A-Z0-9_$#-]+\s+[A-Z0-9_$#-]{1,5}", value):
            status = "ALIAS_EM_NOME_TABELA"
            observation = "A entrada aparenta combinar nome físico e alias SQL na mesma célula."
            action = "Remover o alias do mapping e manter somente o nome físico canônico."
        else:
            status = "ROTULO_NAO_CANONICO"
            observation = "A entrada é texto livre/rótulo e não atende ao formato de nome físico de tabela."
            action = "Informar a tabela física correspondente ou classificar explicitamente como artefato não tabular."
        rows.append({
            "tabela_mapeada": value, "camada_mapeada": rejected["camada"],
            "dominio": rejected.get("dominio", ""), "alertas": "ENTRADA_NAO_CANONICA",
            "sistema_origem": _display_sources({rejected.get("sistema_origem", "")}) or "Não identificado",
            "evidencia_sistema_origem": (
                "Mapping do cliente (coluna Sistema)" if rejected.get("sistema_origem")
                else "Sem Sistema no mapping ou dependência upstream qualificada"
            ),
            "status_fim_a_fim": status, "chega_dashboard": "Não",
            "etapa_alcancada": "Mapping", "qtd_dashboards": 0,
            "dashboards_amostra": "", "tabela_endpoint": "", "caminho_exemplo": value,
            "workspace_id": "", "workspace": "", "report_id": "", "dashboard": "",
            "dataset_id": "", "dataset": "", "tipo_evidencia_endpoint": "",
            "confianca": "", "codigo_motivo": status,
            "observacao": observation, "acao_recomendada": action,
        })
    return rows


def build_dashboard_lineage(
    workspaces_input: str = "input/Workspaces",
    scan_input: str = "input/scan",
    sharedpoint_input: str = "input/sharedpoint",
) -> dict:
    """Monta a linhagem dashboard -> dataset -> referência de tabela.

    Cada referência recebe proveniência e confiança. Nomes vindos apenas da
    entidade do modelo semântico continuam disponíveis para investigação,
    mas não são apresentados como prova de existência física.
    """
    datasets_by_id: dict[str, dict] = {}
    dataflows_by_id: dict[str, dict] = {}
    reports: list[dict] = []
    sharepoint_sources: dict[tuple[str, str, str, str], dict] = {}

    for jf in _load_json_files(workspaces_input):
        data = _read_json(jf)
        if not data:
            continue
        for ws in data.get("workspaces", []):
            ws_id = str(ws.get("id", ""))
            ws_name = ws.get("name", ws_id)
            for ds in ws.get("datasets", []):
                ds_id = str(ds.get("id", ""))
                tables = [
                    _norm(t.get("name", "")) for t in ds.get("tables", [])
                    if t.get("name") and _is_physical_table_name(t.get("name", ""))
                ]
                source_tables: set[str] = set()
                for table in ds.get("tables", []):
                    for source in table.get("source", []):
                        if isinstance(source, dict) and source.get("expression"):
                            source_tables.update(_extract_tables_from_sql(str(source["expression"])))
                for usage in ds.get("datasourceUsages", []):
                    instance = next(
                        (item for item in data.get("datasourceInstances", [])
                         if item.get("datasourceId") == usage.get("datasourceInstanceId")),
                        None,
                    )
                    if not instance:
                        continue
                    details = instance.get("connectionDetails", {})
                    source_url = details.get("sharePointSiteUrl") or details.get("url")
                    if not source_url or "sharepoint" not in str(source_url).lower():
                        continue
                    source_key = (ws_name, ds_id, instance.get("datasourceType", ""), str(source_url))
                    sharepoint_sources[source_key] = {
                        "workspace_id": ws_id, "workspace": ws_name, "dataset_id": ds_id,
                        "dataset": ds.get("name", ds_id),
                        "tipo": instance.get("datasourceType", ""), "fonte": str(source_url),
                        "observacao": "Fonte SharePoint carregada por ingestão para a camada Bronze",
                    }
                dataflow_ids = [
                    str(u.get("targetDataflowId", "")) for u in ds.get("upstreamDataflows", [])
                    if u.get("targetDataflowId")
                ]
                datasets_by_id[ds_id] = {
                    "id": ds_id, "name": ds.get("name", ds_id),
                    "workspace_id": ws_id, "workspace": ws_name,
                    "tables": tables, "source_tables": source_tables,
                    "dataflow_ids": dataflow_ids,
                }
            for df in ws.get("dataflows", []):
                df_id = str(df.get("objectId") or df.get("id", ""))
                dataflows_by_id[df_id] = {
                    "id": df_id, "name": df.get("name", df_id),
                    "workspace_id": ws_id, "workspace": ws_name,
                }
            for rp in ws.get("reports", []):
                reports.append({
                    "workspace_id": ws_id,
                    "workspace": ws_name, "name": rp.get("name", rp.get("id", "")),
                    "report_id": str(rp.get("id", "")),
                    "dataset_id": str(rp.get("datasetId", "")),
                })

    report_keys_by_dataset = {}
    for report in reports:
        report_keys_by_dataset.setdefault(report["dataset_id"], []).append(report)
    sharepoint_rows = []
    for source in sharepoint_sources.values():
        linked_reports = report_keys_by_dataset.get(source["dataset_id"], [])
        if linked_reports:
            for report in linked_reports:
                sharepoint_rows.append({
                    **source, "report_id": report["report_id"],
                    "dashboard": report["name"],
                })
        else:
            sharepoint_rows.append({**source, "report_id": "", "dashboard": ""})

    # Camada/domínio de apoio: reaproveita o escopo estendido (Excel +
    # scripts .sql/.prc/.tab/.vw/.dsx de input/sharedpoint) já classificado.
    scope = compute_global_scope(sharedpoint_input, extended=True)
    camada_by_table: dict[str, dict] = {}
    for r in scope.get("rows", []):
        camada_by_table.setdefault(_norm(r["tabela"]), {"camada": r["camada"], "dominio": r.get("dominio", "")})
    documented_tables = set(camada_by_table)

    def _classify(table_name: str) -> tuple[str, str]:
        info = camada_by_table.get(table_name)
        if info:
            return info["camada"], info["dominio"]
        return _classify_by_prefix(table_name), ""

    # Grafo tabela -> tabela (dependência) extraído dos scripts de
    # input/sharedpoint, pra rastrear pra trás (Gold -> Silver -> Bronze)
    # a partir de qualquer tabela já resolvida. `camada_by_table` (Excel +
    # tudo já descoberto) alimenta o fallback de "nome do arquivo == tabela
    # conhecida" pros scripts .sql/.prc/.txt sem CREATE/INSERT explícito.
    upstream_by_table = build_table_dependency_graph(sharedpoint_input, known_tables=set(camada_by_table.keys()))
    excel_data = read_excel_tables(sharedpoint_input)
    excel_tables = set().union(*excel_data.get("tables_by_camada", {}).values())
    mapped_upstream = build_mapped_table_dependency_graph(sharedpoint_input, excel_tables)
    for target, sources in mapped_upstream.items():
        upstream_by_table.setdefault(target, set()).update(sources)
    mapped_upstream_rows = [
        {
            "tabela_origem": source, "camada_origem": _classify(source)[0],
            "tabela_destino": target, "camada_destino": _classify(target)[0],
        }
        for target, sources in mapped_upstream.items()
        for source in sources
    ]
    upstream_edges: set[tuple[str, str]] = set()

    # Tabelas físicas por Dataflow: extraídas do código M (FROM/JOIN) dos
    # arquivos de input/scan, casados pelo nome normalizado do Dataflow.
    scan_dataflows = _load_scan_dataflows(scan_input)
    tables_by_scan_key = {key: _extract_tables_from_sql(content) for key, content in scan_dataflows.items()}
    dataflow_tables: dict[str, set[str]] = {}
    for df_id, info in dataflows_by_id.items():
        key = _normalize_dataflow_key(info["name"])
        if key in tables_by_scan_key:
            dataflow_tables[df_id] = tables_by_scan_key[key]

    # Cache por dataset: tabelas rastreadas, proveniência de cada endpoint
    # direto e resumo das fontes consultadas.
    resolved_cache: dict[str, tuple[set[str], dict[str, str], str]] = {}

    def _resolve_dataset_tables(ds_id: str) -> tuple[set[str], dict[str, str], str]:
        if ds_id in resolved_cache:
            return resolved_cache[ds_id]
        ds = datasets_by_id.get(ds_id)
        if not ds:
            resolved_cache[ds_id] = (set(), {}, "dataset não resolvido")
            return resolved_cache[ds_id]
        direct_origins: dict[str, str] = {}
        for df_id in ds["dataflow_ids"]:
            for table in dataflow_tables.get(df_id, set()):
                direct_origins[table] = "via Dataflow (M-query)"
        for table in ds.get("source_tables", set()):
            direct_origins.setdefault(table, "via Dataset M-query")
        for table in ds["tables"]:
            direct_origins.setdefault(
                _norm(table), "direto no dataset (nome do modelo, pode não bater com o físico)"
            )
        if not direct_origins:
            resolved_cache[ds_id] = (set(), {}, "sem fontes identificadas")
            return resolved_cache[ds_id]

        origin_kinds = set(direct_origins.values())
        origem = " + ".join(sorted(origin_kinds))

        # Rastreia pra trás (Bronze/Silver) a partir das tabelas resolvidas
        # (normalmente Gold) via o grafo de dependência dos scripts.
        visited, edges = trace_upstream(set(direct_origins), upstream_by_table)
        upstream_edges.update(edges)
        resolved_cache[ds_id] = (visited, direct_origins, origem)
        return resolved_cache[ds_id]

    def _row_origem(t: str, base_origem: str, direct_origins: dict[str, str]) -> str:
        if t in direct_origins:
            return direct_origins[t]
        return f"rastreio SQL/.vw/.tab (upstream — {base_origem})"

    dashboard_rows: list[dict] = []
    seen_dashboards: set[tuple[str, ...]] = set()
    distinct_tables: set[str] = set()
    candidate_model_tables: set[str] = set()
    mquery_tables: set[str] = set()
    high_confidence_references: set[str] = set()
    documented_model_references: set[str] = set()
    low_confidence_model_candidates: set[str] = set()
    reports_with_tables: set[tuple[str, ...]] = set()
    unresolved_dataset_reports = 0
    reports_without_sources = 0
    for rp in reports:
        ds = datasets_by_id.get(rp["dataset_id"])
        tables, direct_origins, origem = _resolve_dataset_tables(rp["dataset_id"])
        report_key = ((rp["report_id"],) if rp["report_id"] else
                      (rp["workspace_id"], rp["name"], rp["dataset_id"]))
        seen_dashboards.add(report_key)
        dataset_name = ds["name"] if ds else "(não resolvido)"
        if not tables:
            if origem == "dataset não resolvido":
                unresolved_dataset_reports += 1
            else:
                reports_without_sources += 1
            dashboard_rows.append({
                "workspace_id": rp["workspace_id"], "workspace": rp["workspace"],
                "report_id": rp["report_id"], "dashboard": rp["name"],
                "dataset_id": rp["dataset_id"], "dataset": dataset_name,
                "tabela": "", "camada": "", "dominio": "", "origem": origem,
                "tipo_evidencia": "NAO_RESOLVIDA", "confianca": "",
                "tipo_objeto": "",
            })
            continue
        reports_with_tables.add(report_key)
        for t in sorted(tables):
            camada, dominio = _classify(t)
            evidence, confidence, object_kind = _reference_evidence(
                t, direct_origins, documented_tables
            )
            distinct_tables.add(t)
            direct_origin = direct_origins.get(t, "")
            if direct_origin.startswith("direto no dataset"):
                candidate_model_tables.add(t)
            elif "M-query" in direct_origin:
                mquery_tables.add(t)
            if confidence == "Alta":
                high_confidence_references.add(t)
            elif confidence == "Média":
                documented_model_references.add(t)
            else:
                low_confidence_model_candidates.add(t)
            dashboard_rows.append({
                "workspace_id": rp["workspace_id"], "workspace": rp["workspace"],
                "report_id": rp["report_id"], "dashboard": rp["name"],
                "dataset_id": rp["dataset_id"], "dataset": dataset_name,
                "tabela": t, "camada": camada, "dominio": dominio,
                "origem": _row_origem(t, origem, direct_origins),
                "tipo_evidencia": evidence, "confianca": confidence,
                "tipo_objeto": object_kind,
            })

    dataset_dataflow_rows: list[dict] = []
    dataset_dataflow_links: list[dict] = []
    for ds_id, ds in datasets_by_id.items():
        for df_id in ds["dataflow_ids"]:
            info = dataflows_by_id.get(df_id)
            if info:
                dataset_dataflow_links.append({
                    "dataset": ds["name"], "dataflow": info["name"],
                    "workspace": ds["workspace"],
                })
        tables, direct_origins, origem = _resolve_dataset_tables(ds_id)
        if not tables:
            dataset_dataflow_rows.append({
                "workspace": ds["workspace"], "tipo": "Dataset", "nome": ds["name"],
                "tabela": "", "camada": "", "dominio": "", "origem": origem,
            })
            continue
        for t in sorted(tables):
            camada, dominio = _classify(t)
            dataset_dataflow_rows.append({
                "workspace": ds["workspace"], "tipo": "Dataset", "nome": ds["name"],
                "tabela": t, "camada": camada, "dominio": dominio,
                "origem": _row_origem(t, origem, direct_origins),
            })
    for df_id, info in dataflows_by_id.items():
        tables = dataflow_tables.get(df_id, set())
        if not tables:
            dataset_dataflow_rows.append({
                "workspace": info["workspace"], "tipo": "Dataflow", "nome": info["name"],
                "tabela": "", "camada": "", "dominio": "", "origem": "sem tabela extraída da M-query (ou sem export em input/scan)",
            })
            continue
        for t in sorted(tables):
            camada, dominio = _classify(t)
            dataset_dataflow_rows.append({
                "workspace": info["workspace"], "tipo": "Dataflow", "nome": info["name"],
                "tabela": t, "camada": camada, "dominio": dominio, "origem": "M-query (FROM/JOIN)",
            })

    excel_tables_by_layer = {
        camada: set(tables)
        for camada, tables in excel_data.get("tables_by_camada", {}).items()
    }
    client_tables = set().union(*excel_tables_by_layer.values()) if excel_tables_by_layer else set()
    client_layers_by_table: dict[str, set[str]] = {}
    client_domains_by_table: dict[str, set[str]] = {}
    for row in excel_data.get("rows", []):
        client_domains_by_table.setdefault(row["tabela"], set()).add(row.get("dominio", ""))
    for camada, tables in excel_tables_by_layer.items():
        for table in tables:
            client_layers_by_table.setdefault(table, set()).add(camada)
    client_layer_conflicts = [
        {"tabela": table, "camadas": ", ".join(sorted(layers))}
        for table, layers in sorted(client_layers_by_table.items())
        if len(layers) > 1
    ]
    client_layer_by_table = {
        table: camada
        for camada, tables in excel_tables_by_layer.items()
        for table in tables
    }
    combined_tables = distinct_tables | client_tables
    camada_totals = Counter(_classify(t)[0] for t in distinct_tables)
    client_layer_totals = Counter(client_layer_by_table.values())
    combined_layer_totals = Counter(_classify(t)[0] for t in combined_tables)

    upstream_rows = [
        {
            "tabela_origem": src, "camada_origem": _classify(src)[0],
            "tabela_destino": dst, "camada_destino": _classify(dst)[0],
        }
        for src, dst in sorted(upstream_edges)
    ]

    # Cruza o escopo obrigatório do cliente (Excel + dependências .sql/.vw)
    # com as tabelas efetivamente alcançadas por cada dashboard. O vínculo
    # só existe para nomes completos idênticos; não há aproximação por
    # prefixo/sufixo ou semelhança textual.
    dashboard_tables: dict[tuple[str, str, str, str, str, str], set[str]] = {}
    dashboard_endpoints: dict[tuple[str, str, str, str, str, str], dict[str, dict]] = {}
    for row in dashboard_rows:
        key = (
            row.get("workspace_id", ""), row["workspace"],
            row.get("report_id", ""), row["dashboard"],
            row.get("dataset_id", ""), row["dataset"],
        )
        if row["tabela"]:
            dashboard_tables.setdefault(key, set()).add(row["tabela"])
            if not row["origem"].startswith("rastreio "):
                dashboard_endpoints.setdefault(key, {})[row["tabela"]] = {
                    "tipo_evidencia": row.get("tipo_evidencia", ""),
                    "confianca": row.get("confianca", ""),
                }
    mapping_dashboard_rows: list[dict] = []
    for mapped_table in sorted(client_tables):
        mapped_closure, _ = trace_upstream({mapped_table}, mapped_upstream)
        consumers = [
            (key, mapped_closure & tables)
            for key, tables in dashboard_tables.items()
            if mapped_closure & tables
        ]
        if not consumers:
            mapping_dashboard_rows.append({
                "tabela_mapeada": mapped_table,
                "tabela_rastreada": mapped_table,
                    "camada": _classify(mapped_table)[0],
                    "camada_mapeada": _classify(mapped_table)[0],
                "workspace_id": "", "workspace": "", "report_id": "",
                "dashboard": "", "dataset_id": "", "dataset": "",
                "status": "Mapeada pelo cliente — sem dashboard consumidor identificado",
            })
            continue
        for (workspace_id, workspace, report_id, dashboard, dataset_id, dataset), matched_tables in consumers:
            for traced_table in sorted(matched_tables):
                mapping_dashboard_rows.append({
                    "tabela_mapeada": mapped_table,
                    "tabela_rastreada": traced_table,
                    "camada": _classify(traced_table)[0],
                    "camada_mapeada": _classify(mapped_table)[0],
                    "workspace_id": workspace_id, "workspace": workspace,
                    "report_id": report_id, "dashboard": dashboard,
                    "dataset_id": dataset_id, "dataset": dataset,
                    "status": "Conexão confirmada por nome completo na linhagem do dashboard",
                })
    mapped_without_dashboard = {
        row["tabela_mapeada"]
        for row in mapping_dashboard_rows
        if not row["dashboard"]
    }
    dependencies_without_dashboard: set[str] = set()
    for mapped_table in mapped_without_dashboard:
        traced_tables, _ = trace_upstream({mapped_table}, mapped_upstream)
        dependencies_without_dashboard.update(traced_tables - {mapped_table})

    end_to_end_mapping_rows = _build_end_to_end_mapping_rows(
        client_tables=client_tables,
        client_layers_by_table=client_layers_by_table,
        client_domains_by_table=client_domains_by_table,
        dependencies_by_target=upstream_by_table,
        dashboard_endpoints=dashboard_endpoints,
        classify=_classify,
        client_sources_by_table=excel_data.get("source_systems_by_table", {}),
        script_sources_by_table=build_table_source_system_hints(sharedpoint_input),
        rejected_rows=excel_data.get("rejected_rows", []),
    )
    end_to_end_status_totals = Counter(row["status_fim_a_fim"] for row in end_to_end_mapping_rows)
    source_system_totals = Counter(
        row["sistema_origem"] for row in end_to_end_mapping_rows
        if row.get("sistema_origem") and row["sistema_origem"] != "Não identificado"
    )

    quality_by_report: dict[tuple[str, ...], dict] = {}
    for row in dashboard_rows:
        report_key = (
            (row.get("report_id", ""),) if row.get("report_id") else
            (row.get("workspace_id", ""), row["dashboard"], row.get("dataset_id", ""))
        )
        quality = quality_by_report.setdefault(report_key, {
            "workspace_id": row.get("workspace_id", ""), "workspace": row["workspace"],
            "report_id": row.get("report_id", ""), "dashboard": row["dashboard"],
            "dataset_id": row.get("dataset_id", ""), "dataset": row["dataset"],
            "referencias": set(), "alta": set(), "media": set(), "baixa": set(),
            "origens_sem_tabela": set(),
        })
        if row["tabela"]:
            quality["referencias"].add(row["tabela"])
            confidence_key = {"Alta": "alta", "Média": "media", "Baixa": "baixa"}.get(row.get("confianca", ""))
            if confidence_key:
                quality[confidence_key].add(row["tabela"])
        elif row["origem"]:
            quality["origens_sem_tabela"].add(row["origem"])

    dashboard_quality_rows: list[dict] = []
    quality_status_totals: Counter = Counter()
    for quality in quality_by_report.values():
        if quality["alta"]:
            status = "RASTREADO_COM_EVIDENCIA_ESTATICA"
        elif quality["media"]:
            status = "CORROBORADO_PELO_INVENTARIO"
        elif quality["baixa"]:
            status = "CANDIDATO_BAIXA_CONFIANCA"
        elif "dataset não resolvido" in quality["origens_sem_tabela"]:
            status = "DATASET_NAO_RESOLVIDO"
        else:
            status = "SEM_FONTE_IDENTIFICADA"
        quality_status_totals[status] += 1
        dashboard_quality_rows.append({
            "workspace_id": quality["workspace_id"], "workspace": quality["workspace"],
            "report_id": quality["report_id"], "dashboard": quality["dashboard"],
            "dataset_id": quality["dataset_id"], "dataset": quality["dataset"],
            "status": status, "total_referencias": len(quality["referencias"]),
            "referencias_alta": len(quality["alta"]),
            "referencias_media": len(quality["media"]),
            "referencias_baixa": len(quality["baixa"]),
            "observacao": "; ".join(sorted(quality["origens_sem_tabela"])),
        })

    return {
        "dashboard_rows": dashboard_rows,
        "dataset_dataflow_rows": dataset_dataflow_rows,
        "dataset_dataflow_links": dataset_dataflow_links,
        "upstream_rows": upstream_rows,
        "mapped_upstream_rows": sorted(
            mapped_upstream_rows,
            key=lambda row: (row["tabela_destino"], row["tabela_origem"]),
        ),
        "mapping_dashboard_rows": mapping_dashboard_rows,
        "end_to_end_mapping_rows": end_to_end_mapping_rows,
        "client_layer_conflicts": client_layer_conflicts,
        "dashboard_quality_rows": sorted(
            dashboard_quality_rows,
            key=lambda row: (row["workspace"], row["dashboard"], row["report_id"]),
        ),
        "sharepoint_rows": sorted(
            sharepoint_rows,
            key=lambda row: (row["workspace"], row["dashboard"], row["dataset"], row["fonte"]),
        ),
        "summary": {
            "total_dashboards": len(seen_dashboards),
            "total_reports_scan": len(reports),
            "total_reports_com_tabelas": len(reports_with_tables),
            "total_reports_dataset_nao_resolvido": unresolved_dataset_reports,
            "total_reports_sem_fontes": reports_without_sources,
            "total_dataflows_scan": len(dataflows_by_id),
            "total_exports_dataflow": len(scan_dataflows),
            "total_dataflows_com_export": len(dataflow_tables),
            "total_tabelas_candidatas_modelo": len(candidate_model_tables),
            "total_tabelas_diretas_mquery": len(mquery_tables),
            "total_referencias_alta_confianca": len(high_confidence_references),
            "total_referencias_modelo_documentadas": len(
                documented_model_references - high_confidence_references
            ),
            "total_candidatas_modelo_baixa_confianca": len(
                low_confidence_model_candidates
                - high_confidence_references
                - documented_model_references
            ),
            "total_tabelas_distintas": len(distinct_tables),
            "tabelas_por_camada": dict(camada_totals),
            "total_tabelas_excel_cliente": len(client_tables),
            "tabelas_por_camada_excel_cliente": dict(client_layer_totals),
            "total_conflitos_camada_excel_cliente": len(client_layer_conflicts),
            "total_tabelas_consolidado": len(combined_tables),
            "tabelas_por_camada_consolidado": dict(combined_layer_totals),
            "tabelas_em_comum": len(client_tables & distinct_tables),
            "tabelas_apenas_excel_cliente": len(client_tables - distinct_tables),
            "tabelas_apenas_dashboards": len(distinct_tables - client_tables),
            "tabelas_excel_com_dashboard": len({
                row["tabela_mapeada"] for row in mapping_dashboard_rows if row["dashboard"]
            }),
            "tabelas_excel_sem_dashboard": len({
                row["tabela_mapeada"] for row in mapping_dashboard_rows if not row["dashboard"]
            }),
            "dependencias_tabelas_excel_sem_dashboard": len(dependencies_without_dashboard),
            "fim_a_fim_por_status": dict(end_to_end_status_totals),
            "total_fim_a_fim_confirmado": sum(
                end_to_end_status_totals[status] for status in ("COMPLETO", "COMPLETO_DIRETO")
            ),
            "total_fim_a_fim_candidato": end_to_end_status_totals["CANDIDATO_BAIXA_CONFIANCA"],
            "total_fim_a_fim_parcial": end_to_end_status_totals["PARCIAL_SEM_DASHBOARD"],
            "total_fim_a_fim_orfao": end_to_end_status_totals["ORFAO_SEM_ARESTA"],
            "total_fim_a_fim_entrada_nao_canonica": sum(
                total for status, total in end_to_end_status_totals.items()
                if status in {
                    "FONTE_ARQUIVO_NAO_TABELA", "NOME_QUALIFICADO_NAO_NORMALIZADO",
                    "ALIAS_EM_NOME_TABELA", "ROTULO_NAO_CANONICO",
                }
            ),
            "total_tabelas_com_sistema_origem": sum(source_system_totals.values()),
            "sistemas_origem": dict(source_system_totals),
            "total_fontes_sharepoint_bronze": len({row["fonte"] for row in sharepoint_rows}),
            "total_dashboards_com_sharepoint": len({
                ((row.get("report_id"),) if row.get("report_id") else
                 (row.get("workspace_id", ""), row["dashboard"], row.get("dataset_id", "")))
                for row in sharepoint_rows if row["dashboard"]
            }),
            "qualidade_por_status": dict(quality_status_totals),
        },
    }


def export_dashboard_lineage_report(result: dict, output_dir: Path, batch_id: str) -> dict[str, str]:
    """Exporta a linhagem para `manifests/dashboard/`: Excel com abas Resumo
    (totais), 'Linhagem Dashboards' (dashboard -> dataset -> tabela) e
    'Dataset e Dataflows' (agregado por dataset/dataflow, independente de
    quantos dashboards os usam)."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    excel_path = output_dir / f"linhagem_dashboards_{batch_id}.xlsx"
    html_path = output_dir / f"linhagem_dashboards_{batch_id}.html"

    wb = openpyxl.Workbook()
    ws_resumo = wb.active
    ws_resumo.title = "Resumo"
    ws_resumo.append(["Métrica", "Valor"])
    summary = result["summary"]
    ws_resumo.append(["Total de dashboards/relatórios", summary["total_dashboards"]])
    ws_resumo.append(["Relatórios no scan (IDs canônicos)", summary.get("total_reports_scan", summary["total_dashboards"])])
    ws_resumo.append(["Relatórios com alguma tabela rastreada", summary.get("total_reports_com_tabelas", 0)])
    ws_resumo.append(["Relatórios com dataset não resolvido", summary.get("total_reports_dataset_nao_resolvido", 0)])
    ws_resumo.append(["Relatórios sem fontes identificadas", summary.get("total_reports_sem_fontes", 0)])
    ws_resumo.append(["Dataflows no scan de workspaces", summary.get("total_dataflows_scan", 0)])
    ws_resumo.append(["Exports de Dataflow fornecidos", summary.get("total_exports_dataflow", 0)])
    ws_resumo.append(["Dataflows casados com export", summary.get("total_dataflows_com_export", 0)])
    ws_resumo.append(["Tabelas candidatas vindas do nome do modelo", summary.get("total_tabelas_candidatas_modelo", 0)])
    ws_resumo.append(["Tabelas físicas diretas extraídas de M-query", summary.get("total_tabelas_diretas_mquery", 0)])
    ws_resumo.append([
        "Referências distintas (inclui nomes candidatos do modelo; não prova existência física)",
        summary["total_tabelas_distintas"],
    ])
    ws_resumo.append(["Referências estáticas via M/SQL/mapping (confiança alta)", summary.get("total_referencias_alta_confianca", 0)])
    ws_resumo.append(["Nomes do modelo corroborados pelo inventário (confiança média)", summary.get("total_referencias_modelo_documentadas", 0)])
    ws_resumo.append(["Nomes somente do modelo semântico (confiança baixa)", summary.get("total_candidatas_modelo_baixa_confianca", 0)])
    for camada, total in sorted(summary["tabelas_por_camada"].items()):
        ws_resumo.append([f"Tabelas na camada {camada}", total])
    ws_resumo.append([])
    ws_resumo.append(["Consolidado — tabelas únicas do Excel + dashboards", summary.get("total_tabelas_consolidado", 0)])
    ws_resumo.append(["Tabelas únicas mapeadas no Excel do cliente", summary.get("total_tabelas_excel_cliente", 0)])
    ws_resumo.append(["Conflitos de camada no Excel do cliente", summary.get("total_conflitos_camada_excel_cliente", 0)])
    ws_resumo.append(["Tabelas únicas alcançadas pelos dashboards", summary.get("total_tabelas_distintas", 0)])
    ws_resumo.append(["Tabelas em comum entre Excel e dashboards", summary.get("tabelas_em_comum", 0)])
    ws_resumo.append(["Apenas no Excel do cliente", summary.get("tabelas_apenas_excel_cliente", 0)])
    ws_resumo.append(["Apenas nos dashboards", summary.get("tabelas_apenas_dashboards", 0)])
    ws_resumo.append(["Tabelas mapeadas pelo cliente sem ligação com dashboards", summary.get("tabelas_excel_sem_dashboard", 0)])
    ws_resumo.append(["Dependências únicas dessas tabelas sem dashboards", summary.get("dependencias_tabelas_excel_sem_dashboard", 0)])
    ws_resumo.append(["Fontes SharePoint na camada Bronze", summary.get("total_fontes_sharepoint_bronze", 0)])
    ws_resumo.append(["Dashboards com fonte SharePoint", summary.get("total_dashboards_com_sharepoint", 0)])
    for camada, total in sorted(summary.get("tabelas_por_camada_consolidado", {}).items()):
        ws_resumo.append([f"Consolidado na camada {camada}", total])
    for camada, total in sorted(summary.get("tabelas_por_camada_excel_cliente", {}).items()):
        ws_resumo.append([f"Excel do cliente na camada {camada}", total])
    for status, total in sorted(summary.get("qualidade_por_status", {}).items()):
        ws_resumo.append([f"Relatórios — {status}", total])
    for status, total in sorted(summary.get("fim_a_fim_por_status", {}).items()):
        ws_resumo.append([f"Mapping fim a fim — {status}", total])
    _autofit(ws_resumo)

    ws_dash = wb.create_sheet("Linhagem Dashboards")
    ws_dash.append([
        "Workspace ID", "Workspace", "Report ID", "Dashboard/Relatório",
        "Dataset ID", "Dataset", "Tabela", "Camada", "Domínio", "Origem",
        "Tipo de evidência", "Confiança", "Tipo de objeto",
    ])
    for r in result["dashboard_rows"]:
        ws_dash.append([
            r.get("workspace_id", ""), r["workspace"], r.get("report_id", ""),
            r["dashboard"], r.get("dataset_id", ""), r["dataset"], r["tabela"],
            r["camada"], r["dominio"], r["origem"],
            r.get("tipo_evidencia", ""), r.get("confianca", ""),
            r.get("tipo_objeto", ""),
        ])
    _autofit(ws_dash)

    ws_dd = wb.create_sheet("Dataset e Dataflows")
    ws_dd.append(["Workspace", "Tipo", "Nome", "Tabela", "Camada", "Domínio", "Origem"])
    for r in result["dataset_dataflow_rows"]:
        ws_dd.append([r["workspace"], r["tipo"], r["nome"], r["tabela"], r["camada"], r["dominio"], r["origem"]])
    _autofit(ws_dd)

    ws_dataset_dataflow = wb.create_sheet("Dataset e Dataflow")
    ws_dataset_dataflow.append(["Workspace", "Dataset", "Dataflow"])
    for r in result.get("dataset_dataflow_links", []):
        ws_dataset_dataflow.append([r["workspace"], r["dataset"], r["dataflow"]])
    _autofit(ws_dataset_dataflow)

    ws_up = wb.create_sheet("Linhagem SQL (upstream)")
    ws_up.append(["Tabela Origem", "Camada Origem", "Tabela Destino", "Camada Destino"])
    for r in result.get("upstream_rows", []):
        ws_up.append([r["tabela_origem"], r["camada_origem"], r["tabela_destino"], r["camada_destino"]])
    _autofit(ws_up)

    ws_mapped = wb.create_sheet("Dependências Mapeamento")
    ws_mapped.append(["Tabela Origem", "Camada Origem", "Tabela Destino", "Camada Destino"])
    for r in result.get("mapped_upstream_rows", []):
        ws_mapped.append([r["tabela_origem"], r["camada_origem"], r["tabela_destino"], r["camada_destino"]])
    _autofit(ws_mapped)

    ws_crosswalk = wb.create_sheet("Mapping até Dashboards")
    ws_crosswalk.append([
        "Tabela Mapeada", "Tabela Rastreada", "Camada", "Camada Mapeada",
        "Workspace ID", "Workspace", "Report ID", "Dashboard/Relatório",
        "Dataset ID", "Dataset", "Status",
    ])
    for r in result.get("mapping_dashboard_rows", []):
        ws_crosswalk.append([
            r["tabela_mapeada"], r["tabela_rastreada"], r["camada"],
            r["camada_mapeada"], r.get("workspace_id", ""), r["workspace"],
            r.get("report_id", ""), r["dashboard"], r.get("dataset_id", ""),
            r["dataset"], r["status"],
        ])
    _autofit(ws_crosswalk)

    ws_e2e = wb.create_sheet("Fim a Fim - Diagnóstico")
    ws_e2e.append([
        "Tabela Mapeada", "Camada Mapeada", "Domínio", "Sistema de Origem",
        "Evidência Sistema de Origem", "Status Fim a Fim",
        "Chega a Dashboard", "Etapa Alcançada", "Qtd. Dashboards", "Dashboards (amostra)",
        "Tabela Endpoint", "Caminho Exemplo", "Workspace ID", "Workspace", "Report ID",
        "Dashboard/Relatório", "Dataset ID", "Dataset", "Tipo Evidência Endpoint",
        "Confiança", "Código Motivo", "Observação", "Ação Recomendada", "Alertas",
    ])
    for r in result.get("end_to_end_mapping_rows", []):
        ws_e2e.append([
            r["tabela_mapeada"], r["camada_mapeada"], r["dominio"],
            r.get("sistema_origem", "Não identificado"), r.get("evidencia_sistema_origem", ""),
            r["status_fim_a_fim"],
            r["chega_dashboard"], r["etapa_alcancada"], r["qtd_dashboards"],
            r["dashboards_amostra"], r["tabela_endpoint"], r["caminho_exemplo"],
            r["workspace_id"], r["workspace"], r["report_id"], r["dashboard"],
            r["dataset_id"], r["dataset"], r["tipo_evidencia_endpoint"], r["confianca"],
            r["codigo_motivo"], r["observacao"], r["acao_recomendada"], r["alertas"],
        ])
    _autofit(ws_e2e)

    ws_legend = wb.create_sheet("Legenda Diagnóstico")
    ws_legend.append(["Código", "Interpretação", "Próxima decisão"])
    ws_legend.append([
        "COMPLETO",
        "Existe caminho estático por nome completo entre a tabela mapeada e um endpoint de dashboard.",
        "Validar materialização e teste funcional; a linhagem não prova runtime.",
    ])
    ws_legend.append([
        "COMPLETO_DIRETO",
        "A própria tabela mapeada é um endpoint referenciado diretamente pelo dataset/Dataflow.",
        "Validar existência e consulta no Fabric.",
    ])
    ws_legend.append([
        "CANDIDATO_BAIXA_CONFIANCA",
        "A coincidência termina apenas em nome de entidade do modelo semântico.",
        "Confirmar o objeto físico antes de contabilizar cobertura.",
    ])
    ws_legend.append([
        "PARCIAL_SEM_DASHBOARD",
        "Há dependências downstream, mas nenhuma chega a endpoint dos dashboards disponíveis.",
        "Investigar o próximo elo indicado em Ação Recomendada.",
    ])
    ws_legend.append([
        "ORFAO_SEM_ARESTA",
        "Nenhum artefato fornecido indica consumidor downstream para a tabela.",
        "Localizar script/pipeline consumidor ou confirmar fora de escopo/descontinuação.",
    ])
    ws_legend.append([
        "CONFLITO_CAMADA",
        "A mesma tabela foi declarada em mais de uma camada no mapping do cliente.",
        "Definir a camada canônica; o alerta não apaga uma cadeia encontrada.",
    ])
    ws_legend.append([
        "FONTE_ARQUIVO_NAO_TABELA",
        "A célula contém um arquivo de origem, não uma tabela física.",
        "Relacionar o arquivo à ingestão Bronze e à tabela materializada.",
    ])
    ws_legend.append([
        "NOME_QUALIFICADO_NAO_NORMALIZADO",
        "A célula contém schema.objeto.",
        "Separar schema e nome físico conforme o contrato do Lakehouse.",
    ])
    ws_legend.append([
        "ALIAS_EM_NOME_TABELA",
        "A célula aparenta conter nome físico e alias SQL.",
        "Remover o alias e manter somente o nome canônico.",
    ])
    ws_legend.append([
        "ROTULO_NAO_CANONICO",
        "A célula contém texto livre/rótulo.",
        "Informar a tabela física ou marcar explicitamente como item não tabular.",
    ])
    ws_legend.append([])
    ws_legend.append(["Limite da evidência", "Resultado estático dos arquivos fornecidos.", "Não prova publicação, execução ou materialização no Fabric."])
    _autofit(ws_legend)

    ws_sharepoint = wb.create_sheet("Fontes SharePoint Bronze")
    ws_sharepoint.append(["Workspace ID", "Workspace", "Report ID", "Dashboard/Relatório", "Dataset ID", "Dataset", "Tipo", "Fonte", "Camada", "Observação"])
    for r in result.get("sharepoint_rows", []):
        ws_sharepoint.append([
            r.get("workspace_id", ""), r["workspace"], r.get("report_id", ""),
            r["dashboard"], r.get("dataset_id", ""), r["dataset"], r["tipo"],
            r["fonte"], "Bronze", r["observacao"],
        ])
    _autofit(ws_sharepoint)

    ws_quality = wb.create_sheet("Qualidade por Dashboard")
    ws_quality.append([
        "Workspace ID", "Workspace", "Report ID", "Dashboard/Relatório",
        "Dataset ID", "Dataset", "Status", "Referências distintas",
        "Alta confiança", "Média confiança", "Baixa confiança", "Observação",
    ])
    for r in result.get("dashboard_quality_rows", []):
        ws_quality.append([
            r["workspace_id"], r["workspace"], r["report_id"], r["dashboard"],
            r["dataset_id"], r["dataset"], r["status"], r["total_referencias"],
            r["referencias_alta"], r["referencias_media"], r["referencias_baixa"],
            r["observacao"],
        ])
    _autofit(ws_quality)

    ws_conflicts = wb.create_sheet("Conflitos de Camada")
    ws_conflicts.append(["Tabela", "Camadas declaradas", "Ação necessária"])
    for r in result.get("client_layer_conflicts", []):
        ws_conflicts.append([
            r["tabela"], r["camadas"],
            "Confirmar camada canônica; nenhuma decisão automática foi tomada",
        ])
    _autofit(ws_conflicts)

    wb.save(str(excel_path))
    export_dashboard_lineage_html(result, html_path)
    return {"excel_path": str(excel_path), "html_path": str(html_path)}


def find_latest_lineage_excel(output_dir: Path) -> Path | None:
    """Encontra o artefato `linhagem_dashboards_*.xlsx` mais recente já
    gerado em `output_dir` (manifests/dashboard/), para a tela carregar sem
    precisar reprocessar do zero a cada acesso."""
    output_dir = Path(output_dir)
    if not output_dir.exists():
        return None
    candidates = sorted(output_dir.glob("linhagem_dashboards_*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    return candidates[0] if candidates else None
