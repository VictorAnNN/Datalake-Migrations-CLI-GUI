"""Motor de linhagem Fabric: extração JSON/ZIP -> Excels + visualização."""
from __future__ import annotations

import base64
import json
import re
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
import plotly.graph_objects as go
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
EVEN_FILL = PatternFill("solid", fgColor="DCE6F1")
ODD_FILL = PatternFill("solid", fgColor="FFFFFF")


@dataclass(frozen=True)
class LineageArtifacts:
    original_excel: Path
    simplified_excel: Path
    detailed_excel: Path
    generated_at: str


def _load_json(path: Path) -> dict[str, Any]:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return json.loads(path.read_text(encoding=enc))
        except Exception:
            continue
    return {}


def _decode_mashup(raw: str) -> str:
    if not raw:
        return ""
    try:
        decoded = base64.b64decode(raw).decode("utf-8-sig", errors="replace")
        if any(k in decoded for k in ("let", "Source", "Oracle", "Sql", "PowerPlatform")):
            return decoded
    except Exception:
        pass
    return raw


def _extract_sql_tables(sql: str) -> list[dict[str, str]]:
    if not sql:
        return []
    sql_clean = sql.replace("#(lf)", "\n").replace("\\n", "\n")
    pattern = (
        r'\b(FROM|JOIN|INNER\s+JOIN|LEFT\s+JOIN|RIGHT\s+JOIN|'
        r'FULL\s+JOIN|CROSS\s+JOIN)\s+"?(\w+)"?\."?(\w+)"?'
    )
    out: list[dict[str, str]] = []
    for match in re.finditer(pattern, sql_clean, re.IGNORECASE):
        out.append({"schema": match.group(2), "table": match.group(3), "usage": match.group(1).split()[0].upper()})
    return out


def _parse_m_expression(expr: str) -> list[dict[str, str]]:
    if not expr or not expr.strip():
        return []
    raw_expr = expr.replace("#(lf)", "\n").replace('\\"', '"')
    sources: list[dict[str, str]] = []

    ws_m = re.search(r'workspaceId\s*=\s*"([^"]+)"', raw_expr)
    df_m = re.search(r'dataflowId\s*=\s*"([^"]+)"', raw_expr)
    ent_m = re.search(r'\[#"([^"]+)"\]', raw_expr) or re.search(r'entity\s*=\s*"([^"]+)"', raw_expr)
    if re.search(r"PowerPlatform\.Dataflows", raw_expr, re.IGNORECASE):
        sources.append(
            {
                "source_type": "Dataflow Gen1",
                "server": "",
                "database": ws_m.group(1) if ws_m else "",
                "schema": "",
                "table": ent_m.group(1) if ent_m else "",
                "url": "",
                "path": "",
                "dataflow_workspace_id": ws_m.group(1) if ws_m else "",
                "dataflow_id": df_m.group(1) if df_m else "",
                "entity": ent_m.group(1) if ent_m else "",
                "usage": "",
                "raw": raw_expr[:400],
            }
        )
        return sources

    for match in re.finditer(
        r'Oracle\.Database\s*\(\s*"([^"]+)"\s*(?:,\s*\[([^\]]*(?:\[[^\]]*\][^\]]*)*)\])?\s*\)',
        raw_expr,
        re.IGNORECASE | re.DOTALL,
    ):
        server = match.group(1)
        opts = match.group(2) or ""
        sql_m = re.search(r'Query\s*=\s*"((?:[^"\\]|\\.|\n)+)"', opts, re.DOTALL) or re.search(
            r'Query\s*=\s*"((?:[^"\\]|\\.|\n)+)"', raw_expr, re.DOTALL
        )
        sql = sql_m.group(1).replace('\\"', '"') if sql_m else ""
        tables = _extract_sql_tables(sql)
        if not tables:
            tables = [{"schema": "", "table": "", "usage": ""}]
        for table in tables:
            sources.append(
                {
                    "source_type": "Oracle",
                    "server": server,
                    "database": "",
                    "schema": table["schema"],
                    "table": table["table"],
                    "url": "",
                    "path": "",
                    "dataflow_workspace_id": "",
                    "dataflow_id": "",
                    "entity": "",
                    "usage": table["usage"],
                    "raw": raw_expr[:400],
                }
            )

    for match in re.finditer(
        r'Sql\.Database(?:s)?\s*\(\s*"([^"]+)"\s*(?:,\s*"([^"]+)")?\s*(?:,\s*\[([^\]]*)\])?\s*\)',
        raw_expr,
        re.IGNORECASE | re.DOTALL,
    ):
        server = match.group(1)
        database = match.group(2) or ""
        opts = match.group(3) or ""
        sql_m = re.search(r'Query\s*=\s*"((?:[^"\\]|\\.)+)"', opts, re.DOTALL)
        sql = sql_m.group(1).replace('\\"', '"') if sql_m else ""
        tables = _extract_sql_tables(sql)
        if not tables:
            tables = [{"schema": "", "table": "", "usage": ""}]
        for table in tables:
            sources.append(
                {
                    "source_type": "SQL Server",
                    "server": server,
                    "database": database,
                    "schema": table["schema"],
                    "table": table["table"],
                    "url": "",
                    "path": "",
                    "dataflow_workspace_id": "",
                    "dataflow_id": "",
                    "entity": "",
                    "usage": table["usage"],
                    "raw": raw_expr[:400],
                }
            )

    for source_type, pattern in (
        ("SharePoint", r'SharePoint\.\w+\s*\(\s*"([^"]+)"'),
        ("Web/API", r'Web\.Contents\s*\(\s*"([^"]+)"'),
        ("OData", r'OData\.Feed\s*\(\s*"([^"]+)"'),
        ("Azure Storage", r'AzureStorage\.\w+\s*\(\s*"([^"]+)"'),
    ):
        for match in re.finditer(pattern, raw_expr, re.IGNORECASE):
            sources.append(
                {
                    "source_type": source_type,
                    "server": "",
                    "database": "",
                    "schema": "",
                    "table": "",
                    "url": match.group(1),
                    "path": "",
                    "dataflow_workspace_id": "",
                    "dataflow_id": "",
                    "entity": "",
                    "usage": "",
                    "raw": raw_expr[:400],
                }
            )

    for match in re.finditer(
        r'(?:Excel\.Workbook|Csv\.Document)\s*\(\s*File\.Contents\s*\(\s*"([^"]+)"', raw_expr, re.IGNORECASE
    ):
        sources.append(
            {
                "source_type": "Excel/CSV/File",
                "server": "",
                "database": "",
                "schema": "",
                "table": "",
                "url": "",
                "path": match.group(1),
                "dataflow_workspace_id": "",
                "dataflow_id": "",
                "entity": "",
                "usage": "",
                "raw": raw_expr[:400],
            }
        )

    if not sources:
        sources.append(
            {
                "source_type": "Unknown/Other",
                "server": "",
                "database": "",
                "schema": "",
                "table": "",
                "url": "",
                "path": "",
                "dataflow_workspace_id": "",
                "dataflow_id": "",
                "entity": "",
                "usage": "",
                "raw": raw_expr[:400],
            }
        )
    return sources


def _classify_dependency(source_type: str) -> str:
    if source_type == "Dataflow Gen1":
        return "Dataflow Gen1 table"
    if source_type in {"Oracle", "SQL Server", "SharePoint", "Web/API", "OData", "Excel/CSV/File", "Azure Storage"}:
        return "Original source"
    return "Unknown/Other"


def _build_paths(dep_type: str, row: dict[str, Any]) -> tuple[str, str]:
    dataset = row.get("dataset_name", "")
    dataset_table = row.get("dataset_table", "")
    if dep_type == "Dataflow Gen1 table":
        src = row.get("source_dataflow_name") or row.get("dataflow_id", "")
        entity = row.get("entity", "")
        as_is = f"Dataflow: {src}{f' / {entity}' if entity else ''} -> Dataset table: {dataset_table} -> Dataset: {dataset}"
        to_be = f"Lake -> Mat. View (ex-Dataflow): {src}{f'.{entity}' if entity else ''} -> Dataset Direct Lake: {dataset}"
        return as_is, to_be
    if dep_type == "Original source":
        src_detail = " > ".join(
            p
            for p in (
                row.get("source_type", ""),
                row.get("server", ""),
                f"{row.get('schema', '')}.{row.get('table', '')}".strip("."),
                row.get("url", ""),
                row.get("path", ""),
            )
            if p
        )
        as_is = f"{src_detail} -> Dataset table: {dataset_table} -> Dataset: {dataset}"
        to_be = f"Lake/Fabric source ({row.get('source_type', '')}) -> Mat. View ({dataset_table}) -> Dataset Direct Lake: {dataset}"
        return as_is, to_be
    return f"Dataset table: {dataset_table} -> Dataset: {dataset}", f"Mat. View: {dataset_table} -> Dataset Direct Lake: {dataset}"


def _autofit_sheet(ws_sheet: Any) -> None:
    for col_idx, col in enumerate(ws_sheet.iter_cols(), start=1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for i, cell in enumerate(col):
            if i == 0:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.fill = EVEN_FILL if i % 2 == 0 else ODD_FILL
            val_len = len(str(cell.value)) if cell.value else 0
            max_len = max(max_len, val_len)
        ws_sheet.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 80)
    ws_sheet.freeze_panes = "A2"


def _write_styled_excel(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.fillna("").to_excel(writer, sheet_name=name, index=False)
    workbook = load_workbook(path)
    for sheet_name in workbook.sheetnames:
        _autofit_sheet(workbook[sheet_name])
    workbook.save(path)


def _extract_data(root: Path) -> dict[str, pd.DataFrame]:
    dataflows_rows: list[dict[str, Any]] = []
    datasets_rows: list[dict[str, Any]] = []
    dataset_tables_rows: list[dict[str, Any]] = []
    lineage_rows: list[dict[str, Any]] = []
    pq_sources_rows: list[dict[str, Any]] = []
    oracle_rows: list[dict[str, Any]] = []
    df_names: dict[str, str] = {}
    datasource_map: dict[str, dict[str, str]] = {}

    for file in root.rglob("*.json"):
        raw = _load_json(file)
        if not raw:
            continue
        for ds_inst in raw.get("datasourceInstances", []):
            ds_id = ds_inst.get("datasourceId", "")
            conn = ds_inst.get("connectionDetails", {})
            datasource_map[ds_id] = {
                "type": ds_inst.get("datasourceType", ""),
                "server": conn.get("server", ""),
                "database": conn.get("database", ""),
                "url": conn.get("url", "") or conn.get("sharePointSiteUrl", ""),
            }

        for ws in raw.get("workspaces", []):
            ws_id = ws.get("id", "unknown")
            ws_name = ws.get("name", ws_id)

            for dataflow in ws.get("dataflows", []):
                dataflow_id = dataflow.get("objectId") or dataflow.get("id", "")
                dataflow_name = dataflow.get("name", dataflow_id)
                df_names[str(dataflow_id).lower()] = dataflow_name
                source_types = set()
                for usage in dataflow.get("datasourceUsages", []):
                    ds_ref = datasource_map.get(usage.get("datasourceInstanceId", ""), {})
                    ds_type = ds_ref.get("type") or "Unknown"
                    source_types.add(ds_type)
                dataflows_rows.append(
                    {
                        "workspace_name": ws_name,
                        "workspace_id": ws_id,
                        "dataflow_name": dataflow_name,
                        "dataflow_id": dataflow_id,
                        "generation": dataflow.get("generation", 1),
                        "datasource_types": "; ".join(sorted(source_types)) if source_types else "",
                    }
                )

            for dataset in ws.get("datasets", []):
                dataset_id = dataset.get("id", "")
                dataset_name = dataset.get("name", dataset_id)
                tables = dataset.get("tables", []) or []
                datasets_rows.append(
                    {
                        "workspace_name": ws_name,
                        "workspace_id": ws_id,
                        "dataset_name": dataset_name,
                        "dataset_id": dataset_id,
                        "table_count": len(tables),
                    }
                )
                for table in tables:
                    table_name = table.get("name", "")
                    dataset_tables_rows.append(
                        {
                            "workspace_name": ws_name,
                            "dataset_name": dataset_name,
                            "dataset_id": dataset_id,
                            "table_name": table_name,
                        }
                    )
                    expression = ""
                    for src in table.get("source", []) or []:
                        expression = _decode_mashup(src.get("expression", ""))
                        if expression:
                            break
                    parsed_sources = _parse_m_expression(expression)
                    for parsed in parsed_sources:
                        dep_type = _classify_dependency(parsed["source_type"])
                        source_dataflow_name = df_names.get(parsed.get("dataflow_id", "").lower(), parsed.get("dataflow_id", ""))
                        row = {
                            "workspace_name": ws_name,
                            "workspace_id": ws_id,
                            "dataset_name": dataset_name,
                            "dataset_id": dataset_id,
                            "dataset_table": table_name,
                            "dependency_type": dep_type,
                            "source_type": parsed.get("source_type", ""),
                            "source_dataflow_name": source_dataflow_name,
                            "dataflow_id": parsed.get("dataflow_id", ""),
                            "dataflow_workspace_id": parsed.get("dataflow_workspace_id", ""),
                            "entity": parsed.get("entity", ""),
                            "server": parsed.get("server", ""),
                            "database": parsed.get("database", ""),
                            "schema": parsed.get("schema", ""),
                            "table": parsed.get("table", ""),
                            "url": parsed.get("url", ""),
                            "path": parsed.get("path", ""),
                            "usage": parsed.get("usage", ""),
                            "power_query_expression": expression[:32000],
                        }
                        as_is, to_be = _build_paths(dep_type, row)
                        row["as_is_path"] = as_is
                        row["to_be_path"] = to_be
                        lineage_rows.append(row)
                        pq_sources_rows.append(
                            {
                                "workspace_name": ws_name,
                                "artifact_name": dataset_name,
                                "table_or_entity": table_name,
                                "source_type": parsed.get("source_type", ""),
                                "server": parsed.get("server", ""),
                                "database": parsed.get("database", ""),
                                "schema": parsed.get("schema", ""),
                                "table": parsed.get("table", ""),
                                "url": parsed.get("url", ""),
                                "path": parsed.get("path", ""),
                                "raw_expression": parsed.get("raw", ""),
                            }
                        )
                        if parsed.get("source_type") == "Oracle" and parsed.get("table"):
                            oracle_rows.append(
                                {
                                    "workspace_name": ws_name,
                                    "artifact_name": dataset_name,
                                    "entity_name": table_name,
                                    "oracle_server": parsed.get("server", ""),
                                    "oracle_schema": parsed.get("schema", ""),
                                    "oracle_table": parsed.get("table", ""),
                                    "usage_context": parsed.get("usage", ""),
                                    "full_qualified": f"{parsed.get('schema', '')}.{parsed.get('table', '')}".strip("."),
                                }
                            )

    return {
        "dataflows": pd.DataFrame(dataflows_rows),
        "datasets": pd.DataFrame(datasets_rows),
        "lineage": pd.DataFrame(lineage_rows),
        "dataset_tables": pd.DataFrame(dataset_tables_rows),
        "pq_sources": pd.DataFrame(pq_sources_rows),
        "oracle_tables": pd.DataFrame(oracle_rows),
    }


def generate_lineage_artifacts(input_path: Path, output_dir: Path) -> LineageArtifacts:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir.mkdir(parents=True, exist_ok=True)
    temp_dir: tempfile.TemporaryDirectory[str] | None = None
    input_root = input_path
    if input_path.suffix.lower() == ".zip":
        temp_dir = tempfile.TemporaryDirectory(prefix="dlctl_lineage_")
        with zipfile.ZipFile(input_path, "r") as archive:
            archive.extractall(temp_dir.name)
        input_root = Path(temp_dir.name)
    if not input_root.exists() or not input_root.is_dir():
        raise ValueError(f"Entrada inválida: {input_path}")

    try:
        data = _extract_data(input_root)
        ws_count = data["datasets"]["workspace_id"].nunique() if not data["datasets"].empty else 0
        readme = pd.DataFrame(
            {
                "Item": [
                    "Gerado em",
                    "Workspaces únicos",
                    "Dataflows Gen1",
                    "Datasets",
                    "Tabelas de dataset",
                    "Linhas de linhagem",
                    "Fontes PQ",
                    "Tabelas Oracle",
                ],
                "Valor": [
                    datetime.now().strftime("%Y-%m-%d %H:%M"),
                    ws_count,
                    len(data["dataflows"]),
                    len(data["datasets"]),
                    len(data["dataset_tables"]),
                    len(data["lineage"]),
                    len(data["pq_sources"]),
                    len(data["oracle_tables"]),
                ],
            }
        )

        original = output_dir / f"fabric_lineage_{timestamp}.xlsx"
        _write_styled_excel(
            original,
            {
                "Read Me": readme,
                "Dataflows Gen1": data["dataflows"],
                "Datasets": data["datasets"],
                "Lineage Simplified": data["lineage"],
                "Dataset Tables": data["dataset_tables"],
                "PQ Sources": data["pq_sources"],
                "Oracle Tables": data["oracle_tables"],
            },
        )

        lineage = data["lineage"].copy()
        def _col(name: str) -> pd.Series:
            if name in lineage:
                return lineage[name].fillna("")
            return pd.Series([""] * len(lineage))

        simplified = output_dir / f"fabric_lineage_simplified_migration_{timestamp}.xlsx"
        source_dataflow_name = _col("source_dataflow_name")
        source_table = _col("table")
        source_url = _col("url")
        source_path = _col("path")
        _write_styled_excel(
            simplified,
            {
                "Lineage Simplified": pd.DataFrame(
                    {
                        "Workspace": _col("workspace_name"),
                        "Workspace ID": _col("workspace_id"),
                        "Dataset": _col("dataset_name"),
                        "Dataset ID": _col("dataset_id"),
                        "Dataset Table": _col("dataset_table"),
                        "Immediate Source Type": _col("source_type"),
                        "Immediate Source": source_dataflow_name.where(source_dataflow_name.astype(str) != "", source_table),
                        "Source Server": _col("server"),
                        "Source Schema": _col("schema"),
                        "Source Table": source_table,
                        "Source Detail": source_url.where(source_url.astype(str) != "", source_path),
                        "AS-IS Path": _col("as_is_path"),
                        "TO-BE Path": _col("to_be_path"),
                        "Power Query Expression (truncated)": _col("power_query_expression").astype(str).str[:2000],
                    }
                ),
                "Dataset Tables": data["dataset_tables"],
                "Dataflows Gen1": data["dataflows"],
                "Read Me": readme,
            },
        )

        detailed = output_dir / f"fabric_lineage_powerquery_detailed_{timestamp}.xlsx"
        source_counts = (
            lineage.groupby(["workspace_name", "dataset_name", "dataset_table"])["source_type"].nunique().reset_index(name="source_type_count")
            if not lineage.empty
            else pd.DataFrame(columns=["workspace_name", "dataset_name", "dataset_table", "source_type_count"])
        )
        _write_styled_excel(
            detailed,
            {
                "Dataflows Gen1": data["dataflows"][["workspace_name", "workspace_id", "dataflow_name", "dataflow_id", "generation"]]
                if not data["dataflows"].empty
                else pd.DataFrame(columns=["workspace_name", "workspace_id", "dataflow_name", "dataflow_id", "generation"]),
                "Datasets": data["datasets"],
                "Lineage Sources": data["pq_sources"],
                "Query Analysis": source_counts,
                "Power Query Steps": pd.DataFrame(
                    {
                        "workspace_name": _col("workspace_name"),
                        "dataset_name": _col("dataset_name"),
                        "dataset_table": _col("dataset_table"),
                        "step_name": "Source",
                        "connector_hint": _col("source_type"),
                        "expression_preview": _col("power_query_expression").astype(str).str[:300],
                    }
                )
                if not lineage.empty
                else pd.DataFrame(
                    columns=["workspace_name", "dataset_name", "dataset_table", "step_name", "connector_hint", "expression_preview"]
                ),
            },
        )
        return LineageArtifacts(
            original_excel=original,
            simplified_excel=simplified,
            detailed_excel=detailed,
            generated_at=timestamp,
        )
    finally:
        if temp_dir is not None:
            temp_dir.cleanup()


def apply_lineage_filters(
    lineage_df: pd.DataFrame,
    workspaces: list[str] | None = None,
    datasets: list[str] | None = None,
    dependency_types: list[str] | None = None,
    source_types: list[str] | None = None,
    include_original_sources: bool = True,
) -> pd.DataFrame:
    out = lineage_df.copy()
    if workspaces:
        out = out[out["workspace_name"].isin(workspaces)]
    if datasets:
        out = out[out["dataset_name"].isin(datasets)]
    if dependency_types:
        out = out[out["dependency_type"].isin(dependency_types)]
    if source_types:
        out = out[out["source_type"].isin(source_types)]
    if not include_original_sources:
        out = out[out["dependency_type"] != "Original source"]
    return out


def build_lineage_sankey_figure(lineage_df: pd.DataFrame, title: str) -> go.Figure:
    if lineage_df.empty:
        fig = go.Figure()
        fig.add_annotation(text="Sem dados para exibir com os filtros atuais.", x=0.5, y=0.5, showarrow=False)
        return fig

    nodes: list[str] = []
    node_index: dict[str, int] = {}
    links_source: list[int] = []
    links_target: list[int] = []
    links_value: list[int] = []

    def _add_node(label: str) -> int:
        if label not in node_index:
            node_index[label] = len(nodes)
            nodes.append(label)
        return node_index[label]

    for _, row in lineage_df.iterrows():
        source_label = (
            row.get("source_dataflow_name")
            if row.get("dependency_type") == "Dataflow Gen1 table"
            else (row.get("table") or row.get("path") or row.get("url") or row.get("source_type") or "Unknown Source")
        )
        source_label = f"SRC: {source_label}"
        dst_table = f"TBL: {row.get('dataset_table') or row.get('dataset_name')}"
        dst_dataset = f"DS: {row.get('dataset_name')}"

        s_idx = _add_node(str(source_label))
        t_idx = _add_node(str(dst_table))
        d_idx = _add_node(str(dst_dataset))

        links_source.extend([s_idx, t_idx])
        links_target.extend([t_idx, d_idx])
        links_value.extend([1, 1])

    fig = go.Figure(
        data=[
            go.Sankey(
                arrangement="snap",
                node=dict(pad=12, thickness=15, label=nodes),
                link=dict(source=links_source, target=links_target, value=links_value),
            )
        ]
    )
    fig.update_layout(title=title, height=700)
    return fig
